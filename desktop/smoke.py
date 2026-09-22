"""Opt-in packaged-app acceptance check using generated data only."""
from io import BytesIO
import json
from pathlib import Path
import time

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from PySide6.QtCore import QTimer
from PySide6.QtWidgets import QApplication

from desktop import VERSION
from v06_model import digest
from v06_patch import current_values


def fixture():
    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "宋体"
    normal.font.size = Pt(12)
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")
    document.add_heading("文序桌面版验证", 0)
    document.add_paragraph("这是一份自动生成的验证文档。调整字体后，文字内容应保持一致。")
    document.add_heading("一、保持文档结构", 1)
    paragraph = document.add_paragraph("不等式示例：")
    math = OxmlElement("m:oMath")
    run = OxmlElement("m:r")
    text = OxmlElement("m:t")
    text.text = "a > b"
    run.append(text)
    math.append(run)
    paragraph._p.append(math)
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "项目", "说明"
    table.cell(1, 0).text, table.cell(1, 1).text = "原格式", "保留表格与公式"
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def start(window, directory):
    root = Path(directory).resolve()
    root.mkdir(parents=True, exist_ok=True)
    # Refuse to overwrite previous acceptance evidence.
    run_dir = root / ("run-" + time.strftime("%Y%m%d-%H%M%S") + "-" + str(time.time_ns())[-6:])
    run_dir.mkdir()
    source = run_dir / "synthetic.docx"
    data = fixture()
    source.write_bytes(data)
    errors = []
    window.error = lambda exc: errors.append(str(exc))
    phase = 0
    started = time.monotonic()
    timer = QTimer(window)
    timer.setInterval(100)
    report = {"version": VERSION, "source": "generated synthetic document", "passed": False}

    def finish(error=None):
        timer.stop()
        report["elapsed_seconds"] = round(time.monotonic() - started, 2)
        if error:
            report["error"] = str(error)
        (run_dir / "result.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        window.confirm_discard = lambda: True
        window.close()
        QApplication.instance().exit(1 if error else 0)

    def advance():
        nonlocal phase
        if window.job:
            return
        try:
            if errors:
                raise RuntimeError(errors[-1])
            if time.monotonic() - started > 600:
                raise RuntimeError("Desktop acceptance timed out")
            if phase == 0:
                phase = 1
                window.open_path(source)
            elif phase == 1:
                assert window.state.session.original == data
                assert window.state.session.versions["original"].data == data
                report["renderer"] = window.state.session.environment["renderer"]
                report["original_hash"] = digest(data)
                report["initial_inventory"] = window.model.inventory
                window.scope.setCurrentIndex(next(i for i in range(window.scope.count()) if window.scope.itemText(i) == "全文文字"))
                window.format_panel.font_combo.setCurrentIndex(window.format_panel.font_combo.findData("黑体"))
                assert window.format_panel.values().get("font_family") == "黑体"
                phase = 2
                window.apply_button.click()
            elif phase == 2:
                assert window.state.session.current != "original"
                assert current_values(window.model, window.model.scope("all"), "font_east_asia") == ["黑体"]
                assert window.model.inventory == report["initial_inventory"]
                current = window.state.session.current
                version = window.state.session.versions[current]
                check = json.loads(version.report_json)
                assert check["page_check"]["differing_pages"]
                window.compare.setChecked(True)
                window.reviewed.setChecked(True)
                target = window.state.save(run_dir / "result.docx")
                assert target.read_bytes() == version.data
                assert source.read_bytes() == data
                window.undo()
                assert window.state.session.current == "original"
                window.choose_version(window.versions.findData(current))
                assert window.reviewed.isChecked()
                report.update(passed=True, output_hash=digest(target.read_bytes()), checks=check["checks"],
                              pages_before=check["pages_before"], pages_after=check["pages_after"])
                phase = 3
            elif phase == 3:
                window.grab().save(str(run_dir / "window.png"))
                finish()
        except Exception as exc:
            finish(exc)

    timer.timeout.connect(advance)
    timer.start()
