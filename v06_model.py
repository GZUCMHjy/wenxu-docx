"""Version-bound OOXML structure; opaque objects never become editing instructions."""
from __future__ import annotations

from dataclasses import dataclass, field
from hashlib import sha256
from io import BytesIO
import posixpath
import re
from urllib.parse import unquote

from docx import Document
from docx.oxml.ns import qn
from lxml import etree

from docx_formatting import FormatError, MAIN, NS, _package, _paragraph_roles, _paragraphs, _xml

MODEL_VERSION = "v06-model-1"
TEXT_TAGS = {qn("w:t"), qn("w:tab"), qn("w:br"), qn("w:cr")}
ROLE_LABELS = {"body": "正文", "title": "文档标题", "heading_1": "一级标题", "heading_2": "二级标题", "heading_3": "三级标题", "label": "栏目标签", "header": "表头", "caption": "图表题注", "unknown": "待归类"}


def digest(data: bytes) -> str:
    return sha256(data).hexdigest()


def node_text(node):
    return "".join((n.text or "") if n.tag == qn("w:t") else "\t" if n.tag == qn("w:tab") else "\n" for n in node.iter() if n.tag in TEXT_TAGS)


def text_runs(paragraph):
    """Only main paragraph text and hyperlinks, never equation/object descendants."""
    return paragraph.xpath("./w:r | ./w:hyperlink/w:r | ./w:fldSimple/w:r", namespaces=NS) if type(paragraph) is etree._Element else paragraph.xpath("./w:r | ./w:hyperlink/w:r | ./w:fldSimple/w:r")


def run_spans(paragraph):
    offset, field_depth = 0, 0
    spans = []
    for run in text_runs(paragraph):
        text = "".join(node_text(c) for c in run if c.tag in TEXT_TAGS)
        marks = run.findall("w:fldChar", NS)
        protected = field_depth > 0 or bool(marks) or run.getparent().tag == qn("w:fldSimple")
        for mark in marks:
            kind = mark.get(qn("w:fldCharType"))
            if kind == "begin":
                field_depth += 1
            elif kind == "end":
                field_depth = max(0, field_depth - 1)
        spans.append((run, offset, offset + len(text), text, protected))
        offset += len(text)
    return spans


def validate_relationships(parts):
    for name, raw in parts.items():
        if not name.endswith(".rels"):
            continue
        root = _xml(raw)
        ids = set()
        base = "" if name == "_rels/.rels" else posixpath.dirname(posixpath.dirname(name))
        for rel in root:
            ident = rel.get("Id")
            if not ident or ident in ids:
                raise FormatError("文档关系 ID 缺失或重复。")
            ids.add(ident)
            if rel.get("TargetMode") == "External":
                continue
            target = unquote(rel.get("Target", "")).split("#", 1)[0]
            path = posixpath.normpath(target.lstrip("/") if target.startswith("/") else posixpath.join(base, target))
            if path not in parts:
                raise FormatError("文档存在缺失的内部关系目标，无法保证对象保全。")
        if name != "_rels/.rels":
            owner = posixpath.join(base, posixpath.basename(name)[:-5])
            if owner in parts and owner.endswith(".xml"):
                for element in _xml(parts[owner]).iter():
                    for key, value in element.attrib.items():
                        if key.startswith("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}") and value not in ids:
                            raise FormatError("文档对象引用了不存在的关系。")


@dataclass
class DocumentModel:
    data: bytes
    revision: str
    source_hash: str
    parts: dict
    root: object
    paragraphs: list[dict]
    cells: list[dict]
    columns: list[dict]
    inventory: dict
    role_overrides: dict = field(default_factory=dict)

    def row(self, pid):
        if type(pid) is not int or not 1 <= pid <= len(self.paragraphs):
            raise FormatError("选区段落不存在。")
        return self.paragraphs[pid - 1]

    def search(self, query):
        if not query:
            return []
        found = []
        for p in self.paragraphs:
            start = 0
            while (start := p["text"].find(query, start)) >= 0:
                end = start + len(query)
                found.append({"kind": "selection", "revision": self.revision, "spans": [{"pid": p["pid"], "start": start, "end": end}], "context": f'{p["location"]} · {p["text"][max(0,start-16):end+16]}'})
                start = end
        return found

    def scope(self, kind, **kwargs):
        return {"kind": kind, "revision": self.revision, **kwargs}


def build_model(data: bytes, revision=None, role_overrides=None):
    parts = _package(data, preservation_profile=True)
    validate_relationships(parts)
    root = _xml(parts[MAIN])
    sections = root.findall(".//w:sectPr", NS)
    if len(sections) != 1:
        raise FormatError("当前仅支持单节稿件；多节尚未完成保全验证。")
    # A field may span paragraphs. Until a cross-paragraph field map is available,
    # protect every paragraph from begin through end (including result text).
    elements = _paragraphs(root)
    roles = _paragraph_roles(root, parts)
    pids = {p: i + 1 for i, p in enumerate(elements)}
    cells, columns, cell_by_element = [], [], {}
    for ti, table in enumerate(root.findall(".//w:tbl", NS), 1):
        grid = table.findall("w:tblGrid/w:gridCol", NS)
        active = {}
        table_cells = []
        for ri, row in enumerate(table.findall("w:tr", NS), 1):
            before = row.find("w:trPr/w:gridBefore", NS)
            col = int(before.get(qn("w:val"), 0)) if before is not None else 0
            next_active = {}
            for ci, cell in enumerate(row.findall("w:tc", NS), 1):
                span = cell.find("w:tcPr/w:gridSpan", NS)
                width = int(span.get(qn("w:val"), 1)) if span is not None else 1
                if width < 1 or (grid and col + width > len(grid)):
                    raise FormatError("表格逻辑网格无效，无法可靠定位合并单元格。")
                vm = cell.find("w:tcPr/w:vMerge", NS)
                continuation = vm is not None and vm.get(qn("w:val"), "continue") == "continue"
                cell_pids = [pids[p] for p in cell.findall(".//w:p", NS) if p in pids and next((a for a in p.iterancestors() if a.tag == qn("w:tc")), None) is cell]
                if continuation:
                    anchor = active.get(col)
                    if not anchor or anchor["end_col"] != col + width:
                        raise FormatError("表格垂直合并链无效。")
                    anchor["end_row"] = ri
                    anchor["pids"].extend(cell_pids)
                    item = anchor
                else:
                    item = {"id": f"t{ti}r{ri}c{col+1}", "table": ti, "row": ri, "end_row": ri, "col": col + 1, "end_col": col + width, "pids": cell_pids, "header": row.find("w:trPr/w:tblHeader", NS) is not None}
                    cells.append(item)
                    table_cells.append(item)
                cell_by_element[cell] = item
                if vm is not None:
                    next_active[col] = item
                col += width
            active = next_active
        for item in table_cells:
            item["text"] = "\n".join(node_text(elements[pid-1]) for pid in item["pids"])
    rows, depth = [], 0
    first_nonempty = next((i for i, p in enumerate(elements, 1) if node_text(p).strip()), None)
    for i, (p, role) in enumerate(zip(elements, roles), 1):
        spans = run_spans(p)
        text = "".join(s[3] for s in spans)
        tc = next((a for a in p.iterancestors() if a.tag == qn("w:tc")), None)
        cell = cell_by_element.get(tc)
        inferred, evidence = role["scope"], role["source"]
        if inferred.startswith("heading_") and inferred not in ROLE_LABELS:
            inferred = "unknown"
        if cell and cell["header"]:
            inferred, evidence = "header", "重复表头行"
        elif inferred == "body" and re.match(r"^(?:[一二三四五六七八九十]+[、．.]|[（(][一二三四五六七八九十]+[）)]|\d+[、．.])", text.strip()):
            inferred, evidence = "unknown", "编号标题候选，需核对角色"
        elif inferred == "body" and re.match(r"^[图表]\s*\d", text.strip()):
            inferred, evidence = "caption", "题注文本候选，可纠正"
        elif inferred == "body" and i == first_nonempty and not cell and len(text.strip()) <= 50:
            align = p.find("w:pPr/w:jc", NS)
            if align is not None and align.get(qn("w:val")) == "center":
                inferred, evidence = "unknown", "首个居中短段落，文档标题候选，需核对角色"
        whole_field = depth > 0
        for mark in p.findall(".//w:fldChar", NS):
            typ = mark.get(qn("w:fldCharType"))
            if typ == "begin":
                depth += 1
            elif typ == "end":
                depth -= 1
            if depth < 0:
                raise FormatError("字段边界无效。")
            whole_field = True
        protected = [(a, b) for _, a, b, _, blocked in spans if blocked and b > a]
        if whole_field:
            protected = [(0, len(text))]
        rows.append({"pid": i, "text": text, "role": inferred, "evidence": evidence, "cell_id": cell["id"] if cell else None, "table": cell["table"] if cell else None, "location": f'表 {cell["table"]} 行 {cell["row"]} 列 {cell["col"]} · 段落 {i}' if cell else f"段落 {i}", "protected_spans": protected, "field": whole_field})
    if depth:
        raise FormatError("字段边界不完整。")
    # Candidates are displayed before adoption, never inferred from comments.
    for cell in cells:
        label = re.sub(r"\s+", "", cell["text"]).strip("：:")
        if not label or len(label) > 24 or not cell["pids"]:
            continue
        same = [c for c in cells if c["table"] == cell["table"]]
        if cell["header"]:
            targets = [c for c in same if c["row"] > cell["end_row"] and not c["header"] and c["col"] >= cell["col"] and c["end_col"] <= cell["end_col"]]
        else:
            targets = [c for c in same if c["row"] == cell["row"] and c["col"] == cell["end_col"] + 1]
        if targets:
            columns.append({"id": cell["id"], "label": label, "table": cell["table"], "pids": [pid for c in targets for pid in c["pids"]], "label_pids": cell["pids"], "evidence": "表头下方填写列" if cell["header"] else "标签右侧填写单元格候选"})
            for pid in cell["pids"]:
                row = rows[pid-1]
                if not cell["header"] and row["role"] == "body" and row["text"].strip():
                    row["role"], row["evidence"] = "unknown", "栏目标签候选，请确认或改为正文"
    overrides = role_overrides or {}
    for pid, role in overrides.items():
        if role not in ROLE_LABELS or type(pid) is not int or not 1 <= pid <= len(rows):
            raise FormatError("角色修正无效。")
        rows[pid-1]["role"], rows[pid-1]["evidence"] = role, "用户修正"
    inventory = {"paragraphs": len(rows), "tables": len(root.findall(".//w:tbl", NS)), "comments": len(_xml(parts["word/comments.xml"])) if "word/comments.xml" in parts else 0, "media_parts": sum(n.startswith("word/media/") for n in parts), "embedded_parts": sum(n.startswith("word/embeddings/") for n in parts), "equations": len(root.xpath(".//*[local-name()='oMath']")), "sections": 1, "capabilities": {"text": "read/preserve/format", "tables": "read/preserve/text-format", "comments_images_equations_fields_embeddings": "preserve-only", "headers_footers": "preserve-only"}}
    return DocumentModel(data, revision or digest(data), digest(data), parts, root, rows, cells, columns, inventory, dict(overrides))
