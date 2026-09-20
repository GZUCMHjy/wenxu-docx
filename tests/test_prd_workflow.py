from copy import deepcopy
from io import BytesIO
import hashlib
import unittest
from zipfile import ZipFile

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.shared import Mm, Pt

from docx_formatting import FormatError
from prd_formatting import (
    _effective_paragraph_value,
    apply_confirmed_rules,
    build_structure_mapping,
    preflight_document,
    safe_stem,
)
from prd_workflow import build_html_report, executable_rules, extract_reference_requirements, normalize_rules, parse_text_requirements, rule_fingerprint


def save(doc):
    out = BytesIO()
    doc.save(out)
    return out.getvalue()


def poc_document():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    doc.styles["Normal"].font.name = "Noto Serif CJK SC"
    doc.styles["Normal"].font.size = Pt(14)
    heading = doc.add_heading("一、教学工作", 1)
    heading.runs[0].font.name = "Noto Serif CJK SC"
    heading.runs[0].font.size = Pt(16)
    doc.add_paragraph("本学期完成了教学任务。")
    doc.add_paragraph("保留这一段。")
    doc.add_paragraph("教务处")
    doc.add_paragraph("2026年9月20日")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "表格内容"
    doc.sections[0].header.paragraphs[0].text = "页眉"
    return save(doc)


def rule(scope, prop, target, unit=""):
    return {"id": f"{scope}-{prop}", "apply": True, "scope": scope, "property": prop, "target": str(target), "unit": unit, "source_type": "测试", "source_detail": "测试规则"}


class RequirementExtractionTests(unittest.TestCase):
    def test_complete_text_splits_objects_and_properties(self):
        rows = parse_text_requirements("一级标题黑体小二号、居中并与下段同页；正文Noto Serif CJK SC四号，固定行距28磅；落款右对齐。")
        pairs = {(row["scope"], row["property"]) for row in rows}
        self.assertTrue({
            ("heading_1", "font_family"), ("heading_1", "font_size"), ("heading_1", "alignment"),
            ("heading_1", "keep_with_next"), ("body", "font_family"), ("body", "font_size"),
            ("body", "line_spacing"), ("signature", "alignment"),
        } <= pairs)

    def test_missing_object_and_unit_are_visible_blockers(self):
        rows = normalize_rules(parse_text_requirements("行距30"))
        selected, blockers = executable_rules(rows)
        self.assertEqual(len(selected), 1)
        self.assertEqual(len(blockers), 1)
        self.assertEqual(blockers[0]["status"], "待补充")
        self.assertIn("作用对象", blockers[0]["reason"])

    def test_conflict_requires_user_resolution(self):
        rows = normalize_rules([rule("body", "font_size", 12, "磅"), rule("body", "font_size", 14, "磅")])
        self.assertEqual([row["status"] for row in rows], ["冲突", "冲突"])
        self.assertEqual(len(executable_rules(rows)[1]), 2)

    def test_unparsed_requirement_is_not_silently_dropped(self):
        rows = parse_text_requirements("正文内容润色得更生动")
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["property"], "unsupported")
        self.assertEqual(rows[0]["status"], "待补充")

    def test_scope_property_mismatches_are_blocked(self):
        rows = normalize_rules([
            rule("page", "font_size", 12, "磅"),
            rule("body", "left_margin", 72, "磅"),
        ])
        self.assertEqual([row["status"] for row in rows], ["待补充", "待补充"])
        self.assertIn("页面对象", rows[0]["reason"])
        self.assertIn("只能作用于页面", rows[1]["reason"])


class PrdFormattingTests(unittest.TestCase):
    def setUp(self):
        self.source = poc_document()
        self.mapping = build_structure_mapping(self.source)
        for row in self.mapping:
            if row["text"] in {"教务处", "2026年9月20日"}:
                row["scope"] = "signature"
            if row["text"] == "保留这一段。":
                row["scope"] = "preserve"
                row["excluded_reason"] = "用户确认保留"

    def test_preflight_inventory_and_scope(self):
        info = preflight_document(self.source)
        self.assertEqual(info["sections"], 1)
        self.assertEqual(info["tables"], 1)
        self.assertEqual(info["headers"], 1)
        for mutate in ("sections", "landscape"):
            doc = Document(BytesIO(self.source))
            if mutate == "sections":
                doc.add_section(WD_SECTION_START.NEW_PAGE)
            else:
                section = doc.sections[0]
                section.orientation = WD_ORIENT.LANDSCAPE
                section.page_width, section.page_height = Mm(297), Mm(210)
            with self.subTest(mutate=mutate), self.assertRaises(FormatError):
                preflight_document(save(doc))

    def test_multi_object_snapshot_preserves_content_and_parts(self):
        raw = [
            rule("heading_1", "font_family", "Noto Sans CJK SC", "字体名"),
            rule("heading_1", "font_size", 18, "磅"),
            rule("heading_1", "alignment", "center", "方式"),
            rule("heading_1", "keep_with_next", "是", "布尔"),
            rule("body", "line_spacing", 28, "固定磅"),
            rule("signature", "alignment", "right", "方式"),
            rule("page", "left_margin", 72, "磅"),
        ]
        rules = normalize_rules(raw, self.source, self.mapping)
        output, summary = apply_confirmed_rules(self.source, rules, self.mapping)
        self.assertTrue(summary["changed"])
        before, after = Document(BytesIO(self.source)), Document(BytesIO(output))
        self.assertEqual(before.element.xpath(".//w:t/text()"), after.element.xpath(".//w:t/text()"))
        self.assertEqual(after.paragraphs[0].runs[0].font.name, "Noto Sans CJK SC")
        self.assertEqual(after.paragraphs[0].runs[0].font.size.pt, 18)
        self.assertEqual(after.paragraphs[0].alignment, WD_ALIGN_PARAGRAPH.CENTER)
        self.assertTrue(_effective_paragraph_value(after.paragraphs[0], "keep_with_next"))
        self.assertEqual(after.paragraphs[1].paragraph_format.line_spacing.pt, 28)
        self.assertIsNone(after.paragraphs[2].paragraph_format.line_spacing)
        self.assertEqual(after.paragraphs[3].alignment, WD_ALIGN_PARAGRAPH.RIGHT)
        self.assertEqual(after.sections[0].left_margin.pt, 72)
        with ZipFile(BytesIO(self.source)) as old, ZipFile(BytesIO(output)) as new:
            for name in old.namelist():
                if name != "word/document.xml":
                    self.assertEqual(old.read(name), new.read(name), name)

    def test_inherited_effective_value_is_noop_and_creates_no_output(self):
        rules = normalize_rules([rule("body", "font_size", 14, "磅")], self.source, self.mapping)
        output, summary = apply_confirmed_rules(self.source, rules, self.mapping)
        self.assertEqual(output, self.source)
        self.assertFalse(summary["changed"])

    def test_second_identical_snapshot_is_byte_identical(self):
        rules = normalize_rules([rule("body", "line_spacing", 28, "固定磅")], self.source, self.mapping)
        first, _ = apply_confirmed_rules(self.source, rules, self.mapping)
        second, summary = apply_confirmed_rules(first, normalize_rules([rule("body", "line_spacing", 28, "固定磅")], first, self.mapping), self.mapping)
        self.assertEqual(first, second)
        self.assertFalse(summary["changed"])

    def test_at_least_line_spacing_is_rewritten_as_exact(self):
        document = Document(BytesIO(self.source))
        paragraph = document.paragraphs[1]
        paragraph.paragraph_format.line_spacing = Pt(28)
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
        source = save(document)
        mapping = build_structure_mapping(source)
        rules = normalize_rules([rule("body", "line_spacing", 28, "固定磅")], source, mapping)
        output, summary = apply_confirmed_rules(source, rules, mapping)
        self.assertTrue(summary["changed"])
        changed = Document(BytesIO(output)).paragraphs[1].paragraph_format
        self.assertEqual(changed.line_spacing.pt, 28)
        self.assertEqual(changed.line_spacing_rule, WD_LINE_SPACING.EXACTLY)

    def test_engine_rejects_incompatible_page_rule(self):
        bad = rule("page", "font_size", 12, "磅")
        bad["normalized_target"] = 12.0
        with self.assertRaisesRegex(FormatError, "页面规则"):
            apply_confirmed_rules(self.source, [bad], self.mapping)

    def test_reference_extracts_styles_without_copying_text(self):
        rows = extract_reference_requirements(self.source)
        self.assertTrue(any(row["scope"] == "heading_1" and row["property"] == "font_size" for row in rows))
        self.assertTrue(all("本学期完成了教学任务" not in row["target"] for row in rows))

    def test_safe_name_and_idempotency_cover_snapshot(self):
        self.assertEqual(safe_stem("../../材料_v12.docx"), ("材料", 13))
        rules = normalize_rules([rule("body", "font_size", 14, "磅")], self.source, self.mapping)
        first = rule_fingerprint(hashlib.sha256(self.source).hexdigest(), rules, self.mapping)
        changed = deepcopy(rules)
        changed[0]["target"] = "15"
        changed = normalize_rules(changed, self.source, self.mapping)
        self.assertNotEqual(first, rule_fingerprint(hashlib.sha256(self.source).hexdigest(), changed, self.mapping))

    def test_report_has_required_delivery_fields(self):
        rules = normalize_rules([rule("body", "font_size", 14, "磅")], self.source, self.mapping)
        report = build_html_report(
            task_id="task-1", input_name="材料.docx", input_hash="abc", output_name="材料_v1.docx",
            parent_version="原稿", rules=rules, mapping=self.mapping,
            summary={"changed_runs": 1, "changed_paragraphs": 0},
            checks=[{"level": "pass", "name": "内容与对象完整性", "detail": "通过"}], environment="Linux test",
        ).decode()
        for text in ("task-1", "abc", "材料_v1.docx", "已应用要求", "检查结果", "未应用规范区域", "Linux test"):
            self.assertIn(text, report)


if __name__ == "__main__":
    unittest.main()
