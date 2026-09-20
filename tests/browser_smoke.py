"""WSL Chromium acceptance: PRD source -> review -> confirmation -> checked artifacts."""

from __future__ import annotations

import hashlib
from io import BytesIO
import json
import os
from pathlib import Path
import platform
import sys
from zipfile import ZipFile

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from playwright.sync_api import expect, sync_playwright

from prd_formatting import _effective_paragraph_value, _effective_run_value


BASE_URL = os.environ.get("WENXU_BASE_URL", "http://127.0.0.1:8501")
REQUIREMENT = "一级标题Noto Sans CJK SC小二号、居中；正文Noto Serif CJK SC四号，固定行距28磅。"


def inspect_output(source: bytes, output: bytes, *, expected_line_spacing=28.0):
    before, after = Document(BytesIO(source)), Document(BytesIO(output))
    assert before.element.xpath(".//w:t/text()") == after.element.xpath(".//w:t/text()")

    heading = next(paragraph for paragraph in after.paragraphs if paragraph.style.name == "Heading 1")
    assert heading.alignment == WD_ALIGN_PARAGRAPH.CENTER
    for run in heading.runs:
        assert _effective_run_value(run, heading, "font_family") == "Noto Sans CJK SC"
        assert _effective_run_value(run, heading, "font_size") == 18

    body = next(paragraph for paragraph in after.paragraphs if paragraph.text == "这是第 1 级标题下的正文。")
    assert _effective_paragraph_value(body, "line_spacing") == {"mode": "exact", "value": expected_line_spacing}
    for run in body.runs:
        assert _effective_run_value(run, body, "font_family") == "Noto Serif CJK SC"
        assert _effective_run_value(run, body, "font_size") == 14

    with ZipFile(BytesIO(source)) as old, ZipFile(BytesIO(output)) as new:
        assert old.namelist() == new.namelist()
        for name in old.namelist():
            if name != "word/document.xml":
                assert old.read(name) == new.read(name), name


def main():
    sample = ROOT / "examples" / "sample.docx"
    source = sample.read_bytes()
    artifacts = ROOT / "verification"
    artifacts.mkdir(exist_ok=True)
    report = {
        "runtime": platform.platform(),
        "python": sys.version.split()[0],
        "source_sha256": hashlib.sha256(source).hexdigest(),
        "checks": [],
    }

    with sync_playwright() as playwright:
        browser = playwright.chromium.launch()
        page = browser.new_page(viewport={"width": 1440, "height": 1000}, accept_downloads=True)
        page.set_default_timeout(120_000)
        page.goto(BASE_URL)
        expect(page.get_by_role("heading", name="文序 · 文档格式纠正")).to_be_visible()

        uploaders = page.locator('input[type="file"]')
        uploaders.nth(0).set_input_files(str(sample))
        expect(page.get_by_text("输入：sample.docx", exact=False)).to_be_visible()
        requirement_box = page.get_by_role("textbox", name="文字要求")
        requirement_box.fill(REQUIREMENT)
        requirement_box.press("Tab")
        extract = page.get_by_role("button", name="提取要求并生成核对表")
        expect(extract).to_be_enabled()
        extract.click()
        expect(page.get_by_text("6 条已选要求完整且无冲突。", exact=True)).to_be_visible()
        expect(page.get_by_text("一级标题Noto Sans CJK SC小二号、居中", exact=True)).to_have_count(3)
        expect(page.get_by_text("正文Noto Serif CJK SC四号", exact=True)).to_have_count(2)
        expect(page.get_by_text("固定行距28磅", exact=True)).to_have_count(1)
        report["checks"].append("A01/A03: one text source expanded into six explicit, conflict-free rules")

        confirm = page.get_by_role("button", name="确认当前要求表与范围")
        execute = page.get_by_role("button", name="确认并执行修改")
        expect(execute).to_be_disabled()
        confirm.click()
        expect(page.get_by_text("已确认，可执行。任何输入、要求或范围变化都会使本确认失效。", exact=True)).to_be_visible()
        expect(execute).to_be_enabled()
        execute.click()

        expect(page.get_by_text("sample_v1.docx 已完成内容、规则、文件和固定环境渲染检查。", exact=True)).to_be_visible(timeout=180_000)
        expect(page.get_by_text("输入预览 · 1 页", exact=True)).to_be_visible()
        expect(page.get_by_text("结果预览 · 1 页", exact=True)).to_be_visible()
        page.screenshot(path=str(artifacts / "browser-prd-result.png"), full_page=True)

        with page.expect_download() as download_info:
            page.get_by_role("button", name="下载修改后的 DOCX").click()
        download = download_info.value
        assert download.suggested_filename == "sample_v1.docx"
        output_path = artifacts / download.suggested_filename
        download.save_as(str(output_path))
        first_output = output_path.read_bytes()
        inspect_output(source, first_output)
        report["checks"].append("F04/F05/F06: checked v1 downloaded; target values, text, and untouched package parts verified")

        with page.expect_download() as report_info:
            page.get_by_role("button", name="下载 HTML 检查报告").click()
        report_path = artifacts / "sample_v1_format_report.html"
        report_info.value.save_as(str(report_path))
        report_html = report_path.read_text(encoding="utf-8")
        for required in ("任务 ID", "输入指纹", "处理环境", "已应用要求", "检查结果", "未应用规范区域"):
            assert required in report_html
        report["checks"].append("F06: independent HTML report contains trace, environment, rules, checks, and preserved-region sections")

        page.get_by_role("button", name="以此结果继续调整").click()
        version_selector = page.get_by_test_id("stSelectbox").filter(has_text="本次输入版本")
        expect(version_selector).to_contain_text("v1 · sample_v1.docx")
        requirement_box = page.get_by_role("textbox", name="文字要求")
        requirement_box.fill("正文固定行距30磅。")
        requirement_box.press("Tab")
        expect(page.get_by_role("button", name="确认并执行修改")).to_be_disabled()
        expect(page.get_by_role("button", name="下载修改后的 DOCX")).to_have_count(0)
        report["checks"].append("A07: editing the source invalidated confirmation and the stale result")

        extract = page.get_by_role("button", name="提取要求并生成核对表")
        expect(extract).to_be_enabled()
        extract.click()
        expect(page.get_by_text("1 条已选要求完整且无冲突。", exact=True)).to_be_visible()
        page.get_by_role("button", name="确认当前要求表与范围").click()
        execute = page.get_by_role("button", name="确认并执行修改")
        expect(execute).to_be_enabled()
        execute.click()
        expect(page.get_by_text("sample_v2.docx 已完成内容、规则、文件和固定环境渲染检查。", exact=True)).to_be_visible(timeout=180_000)
        with page.expect_download() as second_download_info:
            page.get_by_role("button", name="下载修改后的 DOCX").click()
        second_path = artifacts / "sample_v2.docx"
        second_download_info.value.save_as(str(second_path))
        inspect_output(first_output, second_path.read_bytes(), expected_line_spacing=30.0)
        report["checks"].append("A08/A09: v1 used as the explicit input; only body line spacing changed and v2 was created")

        page.set_viewport_size({"width": 390, "height": 844})
        page.screenshot(path=str(artifacts / "browser-prd-mobile.png"), full_page=True)
        browser.close()

    assert sample.read_bytes() == source
    report["checks"].append("source DOCX bytes unchanged")
    report["result"] = "PASS"
    (artifacts / "browser-report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
