import base64
import hashlib
from io import BytesIO
import unittest
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Inches

from docx_formatting import FormatError, apply_formatting, inspect_document, output_filename, validate_result


def sample_docx():
    doc = Document()
    doc.add_heading("教学工作总结", 0)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(9)
    p.paragraph_format.first_line_indent = Pt(24)
    p.paragraph_format.line_spacing = 1.5
    for text, bold, italic in [("本学期", False, False), ("重点工作", True, False), ("与教学回顾。", False, True)]:
        run = p.add_run(text)
        run.bold, run.italic = bold, italic
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor.from_string("222222")
    link = OxmlElement("w:hyperlink")
    rid = doc.part.relate_to("https://example.com", "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    link.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "教学参考"
    run.append(text)
    link.append(run)
    doc.add_paragraph()._p.append(link)
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "项目"
    table.cell(0, 1).text = "完成情况"
    table.cell(1, 0).text = "备课"
    table.cell(1, 1).text = "完成"
    table.cell(1, 1).add_table(rows=1, cols=1).cell(0, 0).text = "嵌套表格"
    table.cell(0, 0).merge(table.cell(0, 1))
    png = base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+aX1sAAAAASUVORK5CYII=")
    doc.add_picture(BytesIO(png), width=Inches(0.25))
    doc.sections[0].header.paragraphs[0].text = "页眉保持原样"
    doc.sections[0].footer.paragraphs[0].text = "页脚保持原样"
    out = BytesIO()
    doc.save(out)
    return out.getvalue()


def replace_part(data, name, transform):
    out = BytesIO()
    with ZipFile(BytesIO(data)) as src, ZipFile(out, "w", ZIP_DEFLATED) as dst:
        for info in src.infolist():
            raw = src.read(info.filename)
            dst.writestr(info, transform(raw) if info.filename == name else raw)
    return out.getvalue()


def heading_sample_docx():
    doc = Document(BytesIO(sample_docx()))
    for level, label in [(1, "一、教学成果"), (2, "（一）课堂教学"), (3, "1. 教学方法")]:
        doc.add_heading(label, level)
        doc.add_paragraph(f"这是第 {level} 级标题下的正文。")
    out = BytesIO()
    doc.save(out)
    return out.getvalue()


class FormattingTests(unittest.TestCase):
    def setUp(self):
        self.source = sample_docx()

    def test_each_option_and_combination_preserve_content_and_other_parts(self):
        options = [
            {"font_size": 14}, {"color": "#C00000"},
            {"line_spacing": {"mode": "exact", "value": 28}},
            {"line_spacing": {"mode": "multiple", "value": 2}},
            {"font_size": 14, "color": "#C00000", "line_spacing": {"mode": "exact", "value": 28}},
        ]
        for option in options:
            with self.subTest(option=option):
                output, summary = apply_formatting(self.source, option)
                self.assertGreater(summary["changed_paragraphs"] + summary["changed_runs"], 0)
                before, after = Document(BytesIO(self.source)), Document(BytesIO(output))
                self.assertEqual(before.element.xpath(".//w:t/text()"), after.element.xpath(".//w:t/text()"))
                p = after.paragraphs[1]
                self.assertEqual(p.paragraph_format.space_after.pt, 9)
                self.assertEqual(p.paragraph_format.first_line_indent.pt, 24)
                self.assertTrue(p.runs[1].bold)
                self.assertTrue(p.runs[2].italic)
                self.assertEqual(p.runs[0].font.size.pt, option.get("font_size", 12))
                self.assertEqual(str(p.runs[0].font.color.rgb), option.get("color", "#222222")[1:])
                with ZipFile(BytesIO(self.source)) as a, ZipFile(BytesIO(output)) as b:
                    self.assertEqual(a.namelist(), b.namelist())
                    for name in a.namelist():
                        if name != "word/document.xml":
                            self.assertEqual(a.read(name), b.read(name), name)
                validate_result(self.source, output, option)

    def test_tables_and_hyperlink_runs_are_formatted_once(self):
        output, _ = apply_formatting(self.source, {"font_size": 16, "color": "#123456"})
        doc = Document(BytesIO(output))
        for run in doc.element.xpath(".//w:body//w:r[w:t]"):
            self.assertEqual(run.find("w:rPr/w:sz", run.nsmap).get(qn("w:val")), "32")
            self.assertEqual(run.find("w:rPr/w:color", run.nsmap).get(qn("w:val")), "123456")

    def test_second_identical_application_is_byte_identical(self):
        once, _ = apply_formatting(self.source, {"color": "#000000"})
        twice, summary = apply_formatting(once, {"color": "#000000"})
        self.assertEqual(once, twice)
        self.assertEqual(summary["changed_runs"], 0)

    def test_theme_color_and_complex_script_size_are_overridden(self):
        doc = Document(BytesIO(self.source))
        rpr = doc.paragraphs[1].runs[0]._r.get_or_add_rPr()
        color = rpr.find(qn("w:color"))
        for attr, value in [("themeColor", "accent1"), ("themeTint", "80"), ("themeShade", "80")]:
            color.set(qn("w:" + attr), value)
        size = OxmlElement("w:szCs")
        size.set(qn("w:val"), "40")
        rpr.find(qn("w:sz")).addnext(size)
        saved = BytesIO()
        doc.save(saved)
        output, _ = apply_formatting(saved.getvalue(), {"color": "#112233", "font_size": 14})
        r = Document(BytesIO(output)).paragraphs[1].runs[0]._r
        self.assertEqual(dict(r.find("w:rPr/w:color", r.nsmap).attrib), {qn("w:val"): "112233"})
        self.assertEqual(r.find("w:rPr/w:szCs", r.nsmap).get(qn("w:val")), "28")

    def test_verifier_rejects_unapplied_targets(self):
        with self.assertRaises(FormatError):
            validate_result(self.source, self.source, {"color": "#112233"})

    def test_source_not_modified(self):
        original_hash = hashlib.sha256(self.source).digest()
        apply_formatting(self.source, {"font_size": 18})
        self.assertEqual(hashlib.sha256(self.source).digest(), original_hash)

    def test_invalid_options(self):
        for options in [{}, {"unknown": 1}, {"color": "red"}, {"font_size": float("nan")}, {"font_size": True}, {"font_size": 0}, {"font_size": 12.25}, {"line_spacing": {"mode": "wrong", "value": 28}}, {"line_spacing": {"mode": "exact", "value": -1}}]:
            with self.subTest(options=options), self.assertRaises(FormatError):
                apply_formatting(self.source, options)

    def test_invalid_documents(self):
        fake = BytesIO()
        with ZipFile(fake, "w") as z:
            z.writestr("document.txt", "not docx")
        for data in [b"", b"%PDF-1.0", b"PKbroken", fake.getvalue(), b"x" * (10 * 1024 * 1024 + 1)]:
            with self.subTest(length=len(data)), self.assertRaises(FormatError):
                apply_formatting(data, {"font_size": 14})

    def test_unverified_structures_and_entities_blocked(self):
        for markup in [b"<w:ins/>", b"<w:txbxContent/>", b"<w:fldSimple/>", b"<w:sdt/>"]:
            bad = replace_part(self.source, "word/document.xml", lambda raw: raw.replace(b"<w:body>", b"<w:body>" + markup))
            with self.subTest(markup=markup), self.assertRaises(FormatError):
                apply_formatting(bad, {"font_size": 14})
        bad = replace_part(self.source, "word/document.xml", lambda raw: raw.replace(b"<w:document", b'<!DOCTYPE x [<!ENTITY x "test">]><w:document', 1))
        with self.assertRaises(FormatError):
            apply_formatting(bad, {"font_size": 14})

    def test_verifier_rejects_text_other_format_and_media_mutation(self):
        options = {"font_size": 14}
        output, _ = apply_formatting(self.source, options)
        for name, transform in [
            ("word/document.xml", lambda raw: raw.replace("本学期".encode(), "被篡改".encode())),
            ("word/document.xml", lambda raw: raw.replace(b'w:after="180"', b'w:after="200"')),
            ("word/media/image1.png", lambda raw: raw + b"corrupt"),
        ]:
            bad = replace_part(output, name, transform)
            with self.subTest(name=name), self.assertRaises(FormatError):
                validate_result(self.source, bad, options)

    def test_safe_versioned_filename(self):
        self.assertEqual(output_filename("总结.docx"), "总结_v1.docx")
        self.assertEqual(output_filename("总结_v2.docx"), "总结_v3.docx")
        self.assertEqual(output_filename("../../总结.docx"), "总结_v1.docx")


class HeadingAndFontTests(unittest.TestCase):
    def test_heading_recognition_style_inheritance_and_explicit_outline(self):
        doc = Document(BytesIO(heading_sample_docx()))
        custom = doc.styles.add_style("学校二级标题", WD_STYLE_TYPE.PARAGRAPH)
        custom.base_style = doc.styles["Heading 2"]
        doc.add_paragraph("继承二级标题", style=custom)
        direct = doc.add_paragraph("直接标记三级大纲")
        outline = OxmlElement("w:outlineLvl")
        outline.set(qn("w:val"), "2")
        direct._p.get_or_add_pPr().append(outline)
        demoted = doc.add_paragraph("覆盖为正文", style="Heading 1")
        outline = OxmlElement("w:outlineLvl")
        outline.set(qn("w:val"), "9")
        demoted._p.get_or_add_pPr().append(outline)
        doc.add_paragraph("看起来像标题的短句").runs[0].bold = True
        doc.add_heading("九级标题", 9)
        out = BytesIO()
        doc.save(out)
        rows = {row["text"]: row for row in inspect_document(out.getvalue())}
        for text, scope in [("教学工作总结", "title"), ("一、教学成果", "heading_1"), ("（一）课堂教学", "heading_2"), ("1. 教学方法", "heading_3"), ("继承二级标题", "heading_2"), ("直接标记三级大纲", "heading_3"), ("覆盖为正文", "body"), ("看起来像标题的短句", "body"), ("九级标题", "heading_9")]:
            self.assertEqual(rows[text]["scope"], scope, text)
        self.assertIn("大纲", rows["直接标记三级大纲"]["source"])

    def test_fonts_and_all_scopes_preserve_untargeted_paragraphs(self):
        source = heading_sample_docx()
        before = Document(BytesIO(source))
        rows = inspect_document(source)
        for scope in ["all", "title", "body", "heading_1", "heading_2", "heading_3"]:
            for family in ["宋体", "黑体"]:
                with self.subTest(scope=scope, family=family):
                    options = {"font_family": family, "font_size": 16, "line_spacing": {"mode": "exact", "value": 30}}
                    output, summary = apply_formatting(source, options, scope=scope)
                    after = Document(BytesIO(output))
                    old_ps = before.element.xpath("./w:body//w:p")
                    new_ps = after.element.xpath("./w:body//w:p")
                    for row, old, new in zip(rows, old_ps, new_ps, strict=True):
                        if scope != "all" and row["scope"] != scope:
                            self.assertEqual(old.xml, new.xml)
                        else:
                            for run in new.xpath(".//w:r[w:t]"):
                                fonts = run.find("w:rPr/w:rFonts", run.nsmap)
                                for attr in ["eastAsia", "ascii", "hAnsi", "cs"]:
                                    self.assertEqual(fonts.get(qn("w:" + attr)), family)
                    self.assertGreater(summary["changed_runs"], 0)
                    validate_result(source, output, options, scope=scope)

    def test_font_themes_removed_hint_preserved_and_second_apply_unchanged(self):
        doc = Document(BytesIO(sample_docx()))
        fonts = doc.paragraphs[1].runs[0]._r.get_or_add_rPr().get_or_add_rFonts()
        fonts.set(qn("w:hint"), "eastAsia")
        for attr in ["asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"]:
            fonts.set(qn("w:" + attr), "minorEastAsia")
        out = BytesIO()
        doc.save(out)
        options = {"font_family": "宋体"}
        result, _ = apply_formatting(out.getvalue(), options)
        fonts = Document(BytesIO(result)).paragraphs[1].runs[0]._r.rPr.rFonts
        self.assertEqual(fonts.get(qn("w:hint")), "eastAsia")
        for attr in ["asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"]:
            self.assertNotIn(qn("w:" + attr), fonts.attrib)
        again, summary = apply_formatting(result, options)
        self.assertEqual(again, result)
        self.assertEqual(summary["changed_runs"], 0)

    def test_scoped_verifier_rejects_global_modification(self):
        source = heading_sample_docx()
        options = {"font_family": "黑体"}
        output, _ = apply_formatting(source, options)
        with self.assertRaises(FormatError):
            validate_result(source, output, options, scope="heading_1")

    def test_invalid_font_scope_and_empty_scope_rejected(self):
        source = heading_sample_docx()
        for font in ["Arial", "", [], None]:
            with self.subTest(font=font), self.assertRaises(FormatError):
                apply_formatting(source, {"font_family": font})
        for scope in ["unknown", "heading_0", [], "heading_7"]:
            with self.subTest(scope=scope), self.assertRaises(FormatError):
                apply_formatting(source, {"font_family": "宋体"}, scope=scope)

    def test_localized_style_without_outline_and_inheritance_cycle(self):
        doc = Document()
        style = doc.styles.add_style("标题 2", WD_STYLE_TYPE.PARAGRAPH)
        doc.add_paragraph("中文样式名称", style=style)
        out = BytesIO()
        doc.save(out)
        self.assertEqual(inspect_document(out.getvalue())[0]["scope"], "heading_2")
        loop = doc.styles.add_style("循环样式", WD_STYLE_TYPE.PARAGRAPH)
        loop.base_style = loop
        doc.add_paragraph("不可解析", style=loop)
        out = BytesIO()
        doc.save(out)
        with self.assertRaises(FormatError):
            inspect_document(out.getvalue())


if __name__ == "__main__":
    unittest.main()
