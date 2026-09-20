"""Deterministic 30-document POC matrix. Run in WSL with WENXU_LIBREOFFICE set."""

from __future__ import annotations

from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
import platform
import sys
from time import perf_counter

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from docx import Document
from docx.shared import Mm, Pt

from prd_formatting import apply_confirmed_rules, build_structure_mapping, current_values, preflight_document
from prd_workflow import normalize_rules, render_docx


KINDS = ("教学总结", "工作通知", "课程方案")
PROFILES = (
    (("body", "line_spacing", "28", "固定磅"),),
    (("body", "font_size", "14", "磅"),),
    (("body", "font_family", "Noto Serif CJK SC", "字体名"),),
    (("heading_1", "font_size", "18", "磅"),),
    (("heading_1", "alignment", "center", "方式"),),
    (("heading_2", "bold", "是", "布尔"),),
    (("signature", "alignment", "right", "方式"),),
    (("page", "left_margin", "72", "磅"),),
    (("body", "first_line_indent", "24", "磅"),),
    (("body", "line_spacing", "1.5", "倍数"), ("heading_1", "font_family", "Noto Sans CJK SC", "字体名")),
)


def document_bytes(kind: str, index: int) -> bytes:
    document = Document()
    section = document.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    document.styles["Normal"].font.name = "Liberation Serif"
    document.styles["Normal"].font.size = Pt(12)
    document.add_heading(f"{kind}验收样本 {index:02d}", 0)
    document.add_heading("一、总体情况", 1)
    document.add_paragraph(f"这是{kind}第 {index} 份样本的正文，包含中文与 English 2026。")
    heading_two = document.add_heading("（一）实施内容", 2)
    for run in heading_two.runs:
        run.font.bold = False
    document.add_paragraph("本段用于验证字号、字体、缩进和行距的属性级修改。")
    document.add_heading("1. 具体安排", 3)
    document.add_paragraph("本段用于验证三级标题之后的正文保持完整。")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "项目"
    table.cell(0, 1).text = "状态"
    table.cell(1, 0).text = f"样本 {index:02d}"
    table.cell(1, 1).text = "完成"
    document.add_paragraph(f"{kind}办公室")
    document.add_paragraph("2026年9月20日")
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def rule(scope: str, prop: str, target: str, unit: str, profile: int):
    return {
        "id": f"p{profile}-{scope}-{prop}",
        "apply": True,
        "scope": scope,
        "property": prop,
        "target": target,
        "unit": unit,
        "source_type": "合成验收金标",
        "source_detail": f"要求变更对 {profile}",
    }


def main():
    started = perf_counter()
    cases = []
    rule_instances = 0
    for number in range(1, 31):
        kind = KINDS[(number - 1) % len(KINDS)]
        profile_number = (number - 1) % len(PROFILES) + 1
        source = document_bytes(kind, number)
        preflight_document(source)
        mapping = build_structure_mapping(source)
        for row in mapping:
            if row["text"] in {f"{kind}办公室", "2026年9月20日"}:
                row["scope"] = "signature"
                row["excluded_reason"] = ""
        rows = [rule(*spec, profile_number) for spec in PROFILES[profile_number - 1]]
        rules = normalize_rules(rows, source, mapping)
        output, summary = apply_confirmed_rules(source, rules, mapping)
        assert summary["changed"], (number, profile_number, summary)
        assert Document(BytesIO(source)).element.xpath(".//w:t/text()") == Document(BytesIO(output)).element.xpath(".//w:t/text()")
        for item in rules:
            values = current_values(output, mapping, item["scope"], item["property"])
            assert values and all(value == item["normalized_target"] for value in values), (number, item, values)
        second, second_summary = apply_confirmed_rules(output, normalize_rules(rows, output, mapping), mapping)
        third, third_summary = apply_confirmed_rules(second, normalize_rules(rows, second, mapping), mapping)
        assert output == second == third
        assert not second_summary["changed"] and not third_summary["changed"]
        rendered = render_docx(output, f"case-{number:02d}")
        assert 1 <= rendered["pages"] <= 50 and rendered["images"]
        rule_instances += len(rules)
        cases.append(
            {
                "case": number,
                "partition": "development" if number <= 20 else "holdout",
                "kind": kind,
                "change_pair": profile_number,
                "rules": len(rules),
                "pages": rendered["pages"],
                "output_sha256": sha256(output).hexdigest(),
                "result": "PASS",
            }
        )

    elapsed = round(perf_counter() - started, 3)
    result = {
        "runtime": platform.platform(),
        "python": platform.python_version(),
        "documents": len(cases),
        "development": 20,
        "holdout": 10,
        "document_types": len(KINDS),
        "requirement_change_pairs": len(PROFILES),
        "rule_instances": rule_instances,
        "content_preservation_rate": 1.0,
        "strict_document_pass_rate": 1.0,
        "idempotent_repetitions": 3,
        "elapsed_seconds": elapsed,
        "cases": cases,
        "result": "PASS",
        "scope_note": "Deterministic synthetic evidence; real pilot documents and Word/WPS acceptance remain separate evidence.",
    }
    artifacts = ROOT / "verification"
    artifacts.mkdir(exist_ok=True)
    (artifacts / "poc-acceptance-30.json").write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({key: value for key, value in result.items() if key != "cases"}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
