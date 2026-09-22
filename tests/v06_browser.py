"""Real local Chromium workflow; generated synthetic data only."""
import json
from io import BytesIO
from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT)); sys.path.insert(0, str(ROOT / 'tests'))
from playwright.sync_api import expect, sync_playwright
from docx import Document
from editor_browser import EditorBrowser
from test_v06 import sample, character_values


def main():
    output = ROOT / 'verification/v06-browser'; output.mkdir(parents=True, exist_ok=True)
    source = sample(complex_objects=False); (output / 'synthetic.docx').write_bytes(source)
    evidence = []
    with sync_playwright() as p:
        browser = p.chromium.launch()
        page = browser.new_page(viewport={'width': 1512, 'height': 1080}, accept_downloads=True)
        ui = EditorBrowser(page); ui.upload(output / 'synthetic.docx')
        def scope_options():
            selector = page.get_by_test_id('stSelectbox').filter(has=page.get_by_text('修改范围', exact=True))
            selector.locator('input').click()
            options = page.get_by_role('option').all_text_contents()
            page.keyboard.press('Escape')
            return options
        assert scope_options() == ['正文', '文档标题', '一级标题', '表格文字', '全文文字（正文区和表格）']
        expect(page.get_by_text('范围详情与排除', exact=True)).to_have_count(0)
        expect(page.get_by_text('本次移除命中段落', exact=True)).to_have_count(0)
        expect(page.locator('[data-testid="stDataFrame"]:visible')).to_have_count(0)
        # A second upload must derive its own choices; empty styled paragraphs
        # neither advertise a heading category nor inflate the text count.
        plain = Document(); plain.add_paragraph('合成普通正文'); plain.add_paragraph(''); plain.add_heading('', 2)
        buffer = BytesIO(); plain.save(buffer); (output / 'plain.docx').write_bytes(buffer.getvalue())
        ui.upload(output / 'plain.docx')
        assert scope_options() == ['正文', '全文文字（正文区和表格）']
        expect(page.get_by_text('已选中 1 段文字', exact=False)).to_be_visible()
        ui.rerun(lambda: page.get_by_text('局部调整', exact=True).click())
        selector = page.get_by_test_id('stSelectbox').filter(has=page.get_by_text('定位方式', exact=True))
        selector.locator('input').click()
        assert page.get_by_role('option').all_text_contents() == ['搜索文字', '段落与字符']
        page.keyboard.press('Escape')
        ui.rerun(lambda: page.get_by_text('全局调整', exact=True).click())
        ui.tab('历史与高级'); ui.expand('核对段落角色与保护区')
        ui.choose('选择需要纠正角色的段落', '段落 1', multi=True, exact=False)
        ui.choose('正确角色', '一级标题')
        ui.click('保存角色修正')
        ui.tab('编辑')
        assert scope_options() == ['一级标题', '全文文字（正文区和表格）']
        expect(page.get_by_text('已选中 1 段文字', exact=False)).to_be_visible()
        ui.upload(output / 'synthetic.docx')
        assert scope_options() == ['正文', '文档标题', '一级标题', '表格文字', '全文文字（正文区和表格）']
        evidence.append('Scope menus follow uploaded structure; empty categories/details are absent; blank paragraphs excluded from text count')
        ui.choose('字号', '五号 · 10.5 磅')
        ui.applied('synthetic_v1.docx')
        ui.confirm()
        v1 = ui.download('下载修改后的 DOCX', output / 'v1.docx')
        assert set(character_values(v1, 2, 'font_size')) == {10.5}
        assert character_values(v1, 2, 'bold') == character_values(source, 2, 'bold')
        evidence.append('A02 one-click font size change; actual download preserves bold')
        ui.click('继续修改')
        ui.rerun(lambda: page.get_by_text('局部调整', exact=True).click())
        ui.text('搜索原文', '重复短语')
        expect(page.get_by_text('找到 2 处；请明确选择一个或多个命中。', exact=True)).to_be_visible()
        ui.choose('选择命中位置', '第 2 处', multi=True, exact=False)
        ui.expand('更多格式：颜色、行距、段落等')
        ui.choose('本次设置的属性（其他保持原样）', '字体颜色', multi=True)
        ui.tab('历史与高级'); ui.expand('核对段落角色与保护区')
        ui.click('将当前选区设为保护区')
        ui.tab('编辑')
        expect(page.get_by_role('button', name='应用并预览', exact=True)).to_be_disabled()
        ui.tab('历史与高级'); ui.expand('核对段落角色与保护区')
        ui.click('解除保护')
        ui.tab('编辑')
        expect(page.get_by_role('button', name='应用并预览', exact=True)).to_be_enabled()
        evidence.append('Selection protection remains available in Advanced; removing it restores the exact selected edit')
        ui.applied('synthetic_v2.docx')
        expect(page.get_by_role('button', name='下载修改后的 DOCX', exact=True)).to_be_disabled()
        ui.confirm()
        v2 = ui.download('下载修改后的 DOCX', output / 'v2.docx')
        assert character_values(v2, 2, 'color') == character_values(v1, 2, 'color')
        assert character_values(v2, 3, 'color')[:4] == ['#FF0000'] * 4
        assert character_values(v2, 2, 'font_size') == character_values(v1, 2, 'font_size')
        ui.expand('检查详情与报告')
        ui.download('下载 HTML 检查报告', output / 'report.html')
        assert 'confirmed' in (output / 'report.html').read_text()
        evidence.append('A03/A09 second search hit only; previous size retained; independent confirmation and report')
        ui.click('撤销本次')
        expect(page.get_by_text('当前结果：synthetic_v1.docx', exact=True)).to_be_visible()
        evidence.append('A13 undo returns preceding version')
        ui.click('继续修改')
        ui.expand('用一句话设置')
        ui.text('一句话修改', '正文小四；正文五号')
        ui.click('解析并加入清单')
        expect(page.get_by_role('button', name='应用并预览', exact=True)).to_be_disabled()
        # Cancelling a conflicting rule must immediately refresh the plan/button.
        ui.expand('第 2 项 · 字号 → 10.5')
        ui.rerun(lambda: page.get_by_text('本次应用', exact=True).nth(1).click())
        expect(page.get_by_role('button', name='应用并预览', exact=True)).to_be_enabled()
        ui.tab('历史与高级'); ui.tab('编辑')
        expect(page.get_by_role('button', name='应用并预览', exact=True)).to_be_enabled()
        # Current live controls are included even with an existing batch.
        ui.choose('字号', '四号 · 14 磅')
        expect(page.get_by_role('button', name='应用并预览', exact=True)).to_be_disabled()
        ui.choose('字号', '保持原样')
        ui.applied('synthetic_v3.docx')
        ui.confirm()
        v3 = ui.download('下载修改后的 DOCX', output / 'v3.docx')
        assert set(character_values(v3, 2, 'font_size')) == {12}
        assert character_values(v3, 3, 'color') == character_values(v1, 3, 'color')
        evidence.append('A11 cancelling conflicts refreshes immediately; tab state survives; live values never ignored; actual saved result matches batch')
        ui.click('继续修改')
        ui.choose('字号', '四号 · 14 磅')
        ui.expand('组合多项修改（可选）')
        ui.click('加入待应用清单')
        expect(page.get_by_role('button', name='应用并预览', exact=True)).to_be_enabled()
        ui.click('清空待应用清单')
        expect(page.get_by_role('button', name='应用并预览', exact=True)).to_be_disabled()
        evidence.append('Optional add-to-batch resets live controls; clearing batch clears pending settings')
        ui.expand('用一句话设置')
        ui.text('一句话修改', '删除首页并发送文件')
        ui.click('解析并加入清单')
        expect(page.get_by_role('button', name='应用并预览', exact=True)).to_be_disabled()
        ui.expand('第 1 项 · 删除首页并发送文件')
        expect(page.get_by_text('未理解：', exact=False)).to_be_visible()
        ui.rerun(lambda: page.get_by_text('本次应用', exact=True).click())
        ui.choose('字号', '四号 · 14 磅')
        expect(page.get_by_role('button', name='应用并预览', exact=True)).to_be_enabled()
        evidence.append('Unparsed text is retained as a blocking cancellable item without crashing')
        page.screenshot(path=str(output / 'editor-final.png'))
        browser.close()
    (output / 'evidence.json').write_text(json.dumps({'checks': evidence}, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({'browser_checks': len(evidence), 'result': 'passed'}))


if __name__ == '__main__': main()
