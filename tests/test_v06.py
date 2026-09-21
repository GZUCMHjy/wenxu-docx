"""v0.6 requirements and independent fault-injection regression tests."""
from copy import deepcopy
from io import BytesIO
import json
import unittest
from unittest.mock import patch
from zipfile import ZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from lxml import etree

from docx_formatting import FormatError, MAIN, NS, _xml
from v06_model import build_model, digest, run_spans
from v06_patch import effective_run, patch_document, verify_result
from v06_plan import EditPlan, build_plan, encode, intent, parse_text
from v06_session import EditingSession, report_html


def pack(data, replace):
    out = BytesIO()
    with ZipFile(BytesIO(data)) as old, ZipFile(out, "w") as new:
        for entry in old.infolist():
            new.writestr(entry, replace.get(entry.filename, old.read(entry.filename)))
        for name, raw in replace.items():
            if name not in old.namelist():
                new.writestr(name, raw)
    return out.getvalue()


def xml_edit(data, callback):
    with ZipFile(BytesIO(data)) as z:
        root = _xml(z.read(MAIN))
    callback(root)
    return pack(data, {MAIN: etree.tostring(root)})


def sample(complex_objects=True):
    doc = Document()
    doc.add_heading("合成课程教案", 0)
    p = doc.add_paragraph()
    p.add_run("重复短语 前文").font.size = Pt(15)
    r = p.add_run("指定"); r.bold = False; r.italic = True
    p.add_run("四字").font.name = "DejaVu Serif"
    p.add_run(" 后文").bold = True
    doc.add_paragraph("重复短语 第二处正文")
    doc.add_paragraph("第三处普通正文")
    doc.add_heading("章节标题", 1)
    table = doc.add_table(rows=3, cols=3)
    table.cell(0, 0).text = "教学目标"
    table.cell(0, 1).merge(table.cell(0, 2)).text = "目标填写内容"
    for c, txt in zip(table.rows[1].cells, ("教学内容", "师生活动", "设计意图")):
        c.text = txt
    table.rows[1]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for c, txt in zip(table.rows[2].cells, ("讲解示例", "练习讨论", "联系生活")):
        c.text = txt
    table2 = doc.add_table(rows=2, cols=2)
    table2.cell(0, 0).merge(table2.cell(1, 0)).text = "纵向合并"
    table2.cell(0, 1).text = "另一格"
    table2.cell(1, 1).text = "下一格"
    out = BytesIO(); doc.save(out)
    data = out.getvalue()
    if not complex_objects:
        return data
    parts = {}
    with ZipFile(BytesIO(data)) as z:
        root = _xml(z.read(MAIN)); rels = _xml(z.read("word/_rels/document.xml.rels")); types = _xml(z.read("[Content_Types].xml"))
    p = root.find("w:body/w:p[2]", NS)
    start = OxmlElement("w:commentRangeStart"); start.set(qn("w:id"), "0")
    end = OxmlElement("w:commentRangeEnd"); end.set(qn("w:id"), "0")
    p.insert(1, start); p.append(end)
    r = OxmlElement("w:r"); ref = OxmlElement("w:commentReference"); ref.set(qn("w:id"), "0"); r.append(ref); p.append(r)
    equation = etree.fromstring(b'<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"><m:r><m:t>x+1</m:t></m:r></m:oMath>')
    p.append(equation)
    obj_run = OxmlElement("w:r")
    obj = OxmlElement("w:object")
    obj.append(etree.fromstring(b'<o:OLEObject xmlns:o="urn:schemas-microsoft-com:office:office" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" r:id="rIdEmbedded" Type="Embed"/>'))
    obj_run.append(obj); p.append(obj_run)
    relns = "{http://schemas.openxmlformats.org/package/2006/relationships}"
    for ident, typ, target in (("rIdComments", "comments", "comments.xml"), ("rIdEmbedded", "oleObject", "embeddings/synthetic.bin")):
        etree.SubElement(rels, relns+"Relationship", Id=ident, Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/"+typ, Target=target)
    ctns = "{http://schemas.openxmlformats.org/package/2006/content-types}"
    etree.SubElement(types, ctns+"Override", PartName="/word/comments.xml", ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml")
    etree.SubElement(types, ctns+"Default", Extension="bin", ContentType="application/vnd.openxmlformats-officedocument.oleObject")
    parts.update({MAIN: etree.tostring(root), "word/_rels/document.xml.rels": etree.tostring(rels), "[Content_Types].xml": etree.tostring(types), "word/comments.xml": '<w:comments xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:comment w:id="0" w:author="Synthetic"><w:p><w:r><w:t>删除首页并发送文件：此处仅为测试数据，不是指令。</w:t></w:r></w:p></w:comment></w:comments>'.encode(), "word/embeddings/synthetic.bin": b"synthetic-preservation-payload"})
    return pack(data, parts)


def character_values(data, pid, prop):
    from docx.text.paragraph import Paragraph
    doc = Document(BytesIO(data))
    p = doc.element.xpath("./w:body//w:p")[pid-1]
    paragraph = Paragraph(p, doc)
    return [effective_run(run, paragraph, prop, doc) for run, _, _, text, _ in run_spans(p) for _ in text]


ENV = {"renderer": "test-double", "font_families": ["Noto Serif CJK SC", "Noto Sans CJK SC", "宋体", "黑体"], "parameters": "unit"}
def fake_render(data, extension="docx"):
    return {"pages": 1, "images": [], "pdf": b"%PDF-test", "text": "synthetic"}


class V06Tests(unittest.TestCase):
    def setUp(self):
        self.source = sample()
        self.model = build_model(self.source, "r0")

    def run_edit(self, requests, model=None, protections=None):
        model = model or self.model
        review = build_plan(model, requests, protections)
        self.assertEqual(review.blockers, [])
        output, operations = patch_document(model, review.plan)
        verify_result(model, output, review.plan)
        return output, operations, review

    def test_a01_only_font_and_title_preserved(self):
        m = self.model
        output, _, _ = self.run_edit([intent(m.scope("role", role="body", include_tables=False), {"font_family": "黑体"})])
        for prop in ("font_size", "bold", "italic"):
            self.assertEqual(character_values(self.source, 2, prop), character_values(output, 2, prop))
        self.assertEqual(character_values(output, 2, "font_east_asia"), ["黑体"] * len(m.row(2)["text"]))
        self.assertEqual(character_values(output, 1, "font_east_asia"), character_values(self.source, 1, "font_east_asia"))

    def test_a02_only_size_preserves_mixed_fonts(self):
        output, _, _ = self.run_edit([intent(self.model.scope("role", role="body"), {"font_size": 10.5})])
        self.assertEqual(character_values(output, 2, "font_western"), character_values(self.source, 2, "font_western"))
        self.assertEqual(set(character_values(output, 2, "font_size")), {10.5})

    def test_a03_second_search_match_only(self):
        hits = self.model.search("重复短语")
        self.assertEqual(len(hits), 2)
        output, _, _ = self.run_edit([intent(hits[1], {"color": "#FF0000"})])
        self.assertEqual(character_values(output, 2, "color"), character_values(self.source, 2, "color"))
        self.assertEqual(character_values(output, 3, "color")[:4], ["#FF0000"] * 4)
        self.assertEqual(character_values(output, 3, "color")[4:], character_values(self.source, 3, "color")[4:])

    def test_a04_four_characters_across_runs_objects_preserved(self):
        hit = self.model.search("指定四字")[0]
        output, _, _ = self.run_edit([intent(hit, {"bold": True})])
        start = hit["spans"][0]["start"]
        before, after = character_values(self.source, 2, "bold"), character_values(output, 2, "bold")
        self.assertEqual(after[start:start+4], [True] * 4)
        self.assertEqual(before[:start], after[:start]); self.assertEqual(before[start+4:], after[start+4:])
        self.assertEqual(build_model(output).inventory, self.model.inventory)

    def test_a05_attribute_merge_order_independent(self):
        m = self.model
        requests = [intent(m.scope("role", role="body"), {"font_family": "宋体", "font_size": 12}), intent(m.scope("paragraphs", pids=[3]), {"font_family": "黑体"}, exception=True)]
        a, _, _ = self.run_edit(requests)
        b, _, _ = self.run_edit(list(reversed(requests)))
        self.assertEqual(a, b)
        self.assertEqual(set(character_values(a, 3, "font_east_asia")), {"黑体"})
        self.assertEqual(set(character_values(a, 3, "font_size")), {12})
        self.assertEqual(set(character_values(a, 2, "font_east_asia")), {"宋体"})

    def test_a06_global_font_local_paragraph_spacing(self):
        m = self.model
        output, _, review = self.run_edit([intent(m.scope("role", role="body"), {"font_family": "宋体"}), intent(m.search("指定四字")[0], {"line_spacing": {"mode": "exact", "value": 28}}, exception=True)])
        self.assertTrue(any("整个段落" in s for s in review.notices))
        self.assertEqual(Document(BytesIO(output)).paragraphs[1].paragraph_format.line_spacing.pt, 28)

    def test_a07_named_filling_cell_and_merges(self):
        col = next(c for c in self.model.columns if c["label"] == "教学目标")
        output, _, _ = self.run_edit([intent(self.model.scope("named_column", label="教学目标"), {"font_family": "黑体", "font_size": 10.5})])
        pid = col["pids"][0]
        self.assertEqual(set(character_values(output, pid, "font_size")), {10.5})
        label = col["label_pids"][0]
        self.assertEqual(character_values(output, label, "font_size"), character_values(self.source, label, "font_size"))
        self.assertEqual(self.model.cells, build_model(output).cells)

    def test_a08_table_body_toggle_and_unknown_notice(self):
        m = self.model
        a = build_plan(m, [intent(m.scope("role", role="body", include_tables=False), {"font_size": 12})])
        b = build_plan(m, [intent(m.scope("role", role="body", include_tables=True), {"font_size": 12})])
        self.assertGreater(len(b.preview), len(a.preview))
        self.assertTrue(b.notices)
        headers = {p["pid"] for p in m.paragraphs if p["role"] == "header"}
        self.assertFalse(headers & {a["pid"] for a in b.preview})

    def test_a09_new_revision_keeps_previous_local_font(self):
        first, _, _ = self.run_edit([intent(self.model.scope("paragraphs", pids=[3]), {"font_family": "黑体"})])
        second = build_model(first, "r1")
        out, _, _ = self.run_edit([intent(second.scope("role", role="body"), {"line_spacing": {"mode": "exact", "value": 28}})], second)
        self.assertEqual(character_values(first, 3, "font_east_asia"), character_values(out, 3, "font_east_asia"))

    def test_a10_global_overrides_history_except_protection(self):
        first, _, _ = self.run_edit([intent(self.model.scope("paragraphs", pids=[2, 3]), {"font_family": "黑体"})])
        m = build_model(first, "r1")
        out, _, review = self.run_edit([intent(m.scope("role", role="body"), {"font_family": "宋体"})], m, [m.scope("paragraphs", pids=[3])])
        self.assertEqual(set(character_values(out, 2, "font_east_asia")), {"宋体"})
        self.assertEqual(set(character_values(out, 3, "font_east_asia")), {"黑体"})
        self.assertTrue(review.excluded)

    def test_a11_conflict_unknown_and_cancel(self):
        m = self.model
        a = intent(m.scope("all"), {"font_size": 12})
        b = intent(m.scope("paragraphs", pids=[3]), {"font_size": 14})
        review = build_plan(m, [a, b])
        self.assertIsNone(review.plan); self.assertTrue(any("冲突" in s for s in review.blockers))
        b["active"] = False
        self.assertIsNotNone(build_plan(m, [a, b]).plan)
        self.assertIsNone(build_plan(m, [intent(m.scope("named_column", label="不存在"), {"bold": True})]).plan)

    def test_a12_complex_objects_and_attachments_preserved(self):
        output, _, _ = self.run_edit([intent(self.model.scope("all"), {"font_size": 12})])
        after = build_model(output)
        for name, data in self.model.parts.items():
            if name != MAIN:
                self.assertEqual(data, after.parts[name])
        self.assertIn("不是指令", after.parts["word/comments.xml"].decode())

    def test_a13_atomic_versions_undo_and_stale_plan(self):
        session = EditingSession.load("synthetic.docx", self.source, renderer=fake_render)
        m = session.model()
        plan = build_plan(m, [intent(m.scope("all"), {"font_size": 12})]).plan
        first, created = session.execute(m, plan, environment=ENV)
        self.assertTrue(created)
        again, created = session.execute(m, plan, environment=ENV)
        self.assertEqual(first.id, again.id); self.assertFalse(created)
        session.select(first.id)
        with self.assertRaises(FormatError): session.execute(m, plan, environment=ENV)
        session.undo(); self.assertEqual(session.current, "original")
        self.assertEqual(session.original, self.source)

    def test_a14_text_and_ui_same_operations(self):
        m = self.model
        text = parse_text(m, "正文宋体小四，但第三段黑体")
        ui = [intent(m.scope("role", role="body", include_tables=True), {"font_family": "宋体", "font_size": 12}), intent(m.scope("paragraphs", pids=[3]), {"font_family": "黑体"}, exception=True)]
        a, ops_a, _ = self.run_edit(text)
        b, ops_b, _ = self.run_edit(ui)
        self.assertEqual(a, b); self.assertEqual(ops_a, ops_b)

    def test_partial_overlap_and_protected_characters(self):
        m = self.model
        scope = lambda a,b: m.scope("selection", spans=[{"pid": 2, "start": a, "end": b}])
        requests = [intent(scope(0, 8), {"color": "#FF0000"}), intent(scope(3, 6), {"color": "#000000"}, exception=True)]
        out, _, _ = self.run_edit(requests, protections=[scope(4,5)])
        values = character_values(out, 2, "color")
        self.assertEqual(values[:3], ["#FF0000"] * 3)
        self.assertEqual(values[3], "#000000"); self.assertEqual(values[5], "#000000")
        self.assertEqual(values[4], character_values(self.source, 2, "color")[4])

    def test_every_control_and_idempotency(self):
        values = {"font_east_asia": "宋体", "font_western": "Noto Serif CJK SC", "font_size": 12, "bold": False, "italic": False, "underline": True, "color": "#00AA99", "alignment": "justify", "first_line_indent": 21, "left_indent": 3, "right_indent": 4, "space_before": 5, "space_after": 6, "line_spacing": {"mode": "multiple", "value": 1.5}, "keep_with_next": True}
        out, _, _ = self.run_edit([intent(self.model.scope("paragraphs", pids=[2]), values), intent(self.model.scope("page"), {"top_margin": 65, "bottom_margin": 66, "left_margin": 67, "right_margin": 68})])
        m = build_model(out, "r1")
        same, operations, _ = self.run_edit([intent(m.scope("paragraphs", pids=[2]), values), intent(m.scope("page"), {"top_margin": 65, "bottom_margin": 66, "left_margin": 67, "right_margin": 68})], m)
        self.assertEqual(same, out); self.assertEqual(operations, [])

    def test_forged_or_stale_plan_is_blocked(self):
        plan = build_plan(self.model, [intent(self.model.scope("all"), {"font_size": 12})]).plan
        data = plan.read(); data["assignments"][0]["value"] = 72
        payload = encode(data)
        with self.assertRaises(FormatError): patch_document(self.model, EditPlan(payload, digest(payload.encode())))
        with self.assertRaises(FormatError): patch_document(build_model(self.source, "new"), plan)

    def test_validator_rejects_content_property_object_and_part_tampering(self):
        out, _, review = self.run_edit([intent(self.model.search("指定四字")[0], {"bold": True})])
        faults = [
            xml_edit(out, lambda root: setattr(root.find(".//w:t", NS), "text", "篡改")),
            xml_edit(out, lambda root: root.find(".//w:object", NS).getparent().remove(root.find(".//w:object", NS))),
            pack(out, {"word/embeddings/synthetic.bin": b"changed"}),
            xml_edit(out, lambda root: root.find(".//w:rPr/w:i", NS).set(qn("w:val"), "0")),
        ]
        for fault in faults:
            with self.subTest(hash=digest(fault)), self.assertRaises(FormatError): verify_result(self.model, fault, review.plan)

    def test_broken_relationship_and_revision_content_rejected(self):
        with self.assertRaises(FormatError):
            build_model(pack(self.source, {"word/_rels/document.xml.rels": self.model.parts["word/_rels/document.xml.rels"].replace(b"embeddings/synthetic.bin", b"embeddings/missing.bin")}))
        def add_revision(root): root.find("w:body/w:p", NS).append(OxmlElement("w:ins"))
        with self.assertRaises(FormatError): build_model(xml_edit(self.source, add_revision))

    def test_unknown_suffix_not_silently_accepted(self):
        for text in ("正文宋体并删除首页", "正文黑体五号并且润色内容", "正文行距28", "教学目标黑体"):
            with self.subTest(text=text): self.assertIsNone(build_plan(self.model, parse_text(self.model, text)).plan)

    def test_missing_fonts_and_renderer_failure_do_not_create_versions(self):
        s = EditingSession.load("test.docx", self.source, renderer=fake_render); m = s.model()
        p = build_plan(m, [intent(m.scope("all"), {"font_family": "Unknown Font"})]).plan
        with self.assertRaises(FormatError): s.execute(m, p, environment=ENV)
        p = build_plan(m, [intent(m.scope("all"), {"font_size": 12})]).plan
        s.renderer = lambda *args: (_ for _ in ()).throw(FormatError("injected render error"))
        with self.assertRaises(FormatError): s.execute(m, p, environment=ENV)
        self.assertEqual(list(s.versions), ["original"]); self.assertEqual(s.next_number, 1)

    def test_noop_does_not_allocate_a_version(self):
        s = EditingSession.load("test_v12.docx", self.source, renderer=fake_render); m = s.model()
        p = build_plan(m, [intent(m.scope("all"), {"font_size": 12})]).plan
        v, _ = s.execute(m, p, environment=ENV); self.assertEqual(v.id, "v13")
        s.select(v.id); m = s.model()
        p = build_plan(m, [intent(m.scope("all"), {"font_size": 12})]).plan
        self.assertEqual(s.execute(m, p, environment=ENV), (None, False))
        self.assertEqual(s.next_number, 14)

    def test_report_escapes_source_and_separates_visual_review(self):
        s = EditingSession.load("test.docx", self.source, renderer=fake_render); m = s.model()
        p = build_plan(m, [intent(m.scope("all"), {"font_size": 12}, source="<script>alert(1)</script>")]).plan
        v, _ = s.execute(m, p, environment=ENV)
        html = report_html(v).decode()
        self.assertNotIn("<script>", html); self.assertIn("&lt;script&gt;", html)
        self.assertIn('"visual_review": "pending"'.replace('"', '&quot;'), html)

    def test_hyperlink_and_multiple_text_nodes(self):
        def change(root):
            p = root.find("w:body/w:p[4]", NS)
            for c in list(p): p.remove(c)
            hyperlink = OxmlElement("w:hyperlink"); hyperlink.set(qn("w:anchor"), "bookmark")
            r = OxmlElement("w:r")
            for txt in ("ABC", "DEF"):
                t = OxmlElement("w:t"); t.text = txt; r.append(t)
            hyperlink.append(r); p.append(hyperlink)
        data = xml_edit(self.source, change); m = build_model(data)
        out, _, _ = self.run_edit([intent(m.search("BCDE")[0], {"underline": True})], m)
        self.assertEqual(character_values(out, 4, "underline"), [False, True, True, True, True, False])

    def test_field_results_are_excluded_and_untouched(self):
        def change(root):
            p = root.find("w:body/w:p[4]", NS)
            r = OxmlElement("w:r"); f = OxmlElement("w:fldChar"); f.set(qn("w:fldCharType"), "begin"); r.append(f); p.insert(0, r)
            r = OxmlElement("w:r"); f = OxmlElement("w:fldChar"); f.set(qn("w:fldCharType"), "end"); r.append(f); p.append(r)
        m = build_model(xml_edit(self.source, change))
        out, _, review = self.run_edit([intent(m.scope("all"), {"font_size": 12})], m)
        self.assertTrue(review.excluded)
        self.assertEqual(character_values(m.data, 4, "font_size"), character_values(out, 4, "font_size"))

    def test_empty_comment_reference_run_inside_selection_is_not_text(self):
        def change(root):
            p = root.find("w:body/w:p[2]", NS)
            run = OxmlElement("w:r")
            ref = OxmlElement("w:commentReference"); ref.set(qn("w:id"), "0")
            run.append(ref); p.insert(3, run)
        m = build_model(xml_edit(self.source, change))
        out, _, _ = self.run_edit([intent(m.scope("all"), {"font_size": 12})], m)
        self.assertEqual(set(character_values(out, 2, "font_size")), {12})

    def test_complex_script_size_is_not_a_false_noop(self):
        def change(root):
            run = root.find("w:body/w:p[3]/w:r", NS)
            rpr = OxmlElement("w:rPr")
            for tag, value in (("sz", "24"), ("szCs", "40")):
                node = OxmlElement("w:"+tag); node.set(qn("w:val"), value); rpr.append(node)
            run.insert(0, rpr)
        m = build_model(xml_edit(self.source, change))
        out, operations, _ = self.run_edit([intent(m.scope("paragraphs", pids=[3]), {"font_size": 12})], m)
        self.assertTrue(operations)
        self.assertEqual(set(character_values(out, 3, "font_size")), {12})

    def test_unquoted_named_column_uses_same_plan(self):
        m = self.model
        out, _, _ = self.run_edit(parse_text(m, "教学目标这一栏改黑体五号"))
        col = next(c for c in m.columns if c["label"] == "教学目标")
        self.assertEqual(set(character_values(out, col["pids"][0], "font_size")), {10.5})

    def test_direct_formatted_title_and_multiline_label_stay_out_of_body(self):
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        doc = Document()
        doc.add_paragraph("合成封面标题").alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_paragraph("普通正文")
        t = doc.add_table(rows=1, cols=2)
        t.cell(0, 0).text = "教学重点"
        t.cell(0, 0).add_paragraph("与难点")
        t.cell(0, 1).text = "填写内容"
        data = BytesIO(); doc.save(data); m = build_model(data.getvalue())
        self.assertEqual(m.row(1)["role"], "unknown")
        column = next(c for c in m.columns if c["label"] == "教学重点与难点")
        self.assertTrue(all(m.row(pid)["role"] == "unknown" for pid in column["label_pids"]))
        self.run_edit([intent(m.scope("named_column", label="教学重点与难点"), {"font_size": 10.5})], m)


if __name__ == "__main__":
    unittest.main()
