"""PRD v0.5 formatting engine: confirmed multi-rule snapshots and strict preservation."""

from __future__ import annotations

from copy import deepcopy
from io import BytesIO
import json
import math
from pathlib import PurePosixPath
import re
from zipfile import ZipFile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.text.paragraph import Paragraph
from docx.text.run import Run
from lxml import etree

from docx_formatting import (
    FONT_ATTRS,
    FONT_THEMES,
    MAIN,
    NS,
    FormatError,
    _package,
    _paragraphs,
    _runs,
    _xml,
    inspect_document,
)


ROLE_LABELS = {
    "all": "全部段落",
    "title": "文档标题/副标题",
    "heading_1": "一级标题",
    "heading_2": "二级标题",
    "heading_3": "三级标题",
    "body": "正文",
    "signature": "落款",
    "preserve": "原样保留",
    "page": "页面",
}

PROPERTY_LABELS = {
    "font_family": "字体",
    "font_size": "字号",
    "bold": "加粗",
    "italic": "斜体",
    "color": "字体颜色",
    "alignment": "对齐",
    "first_line_indent": "首行缩进",
    "left_indent": "左缩进",
    "right_indent": "右缩进",
    "space_before": "段前",
    "space_after": "段后",
    "line_spacing": "行距",
    "keep_with_next": "与下段同页",
    "top_margin": "上页边距",
    "bottom_margin": "下页边距",
    "left_margin": "左页边距",
    "right_margin": "右页边距",
}

CHARACTER_PROPERTIES = {"font_family", "font_size", "bold", "italic", "color"}
PARAGRAPH_PROPERTIES = {
    "alignment",
    "first_line_indent",
    "left_indent",
    "right_indent",
    "space_before",
    "space_after",
    "line_spacing",
    "keep_with_next",
}
PAGE_PROPERTIES = {"top_margin", "bottom_margin", "left_margin", "right_margin"}
SUPPORTED_PROPERTIES = CHARACTER_PROPERTIES | PARAGRAPH_PROPERTIES | PAGE_PROPERTIES

ALIGNMENTS = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}


def preflight_document(data: bytes) -> dict:
    """Validate the declared POC input subset and return a complexity inventory."""
    parts = _package(data)
    try:
        document = Document(BytesIO(data))
    except Exception as exc:
        raise FormatError("无法读取此 DOCX，请检查文件是否完整。") from exc
    if len(document.sections) != 1:
        raise FormatError("POC 当前仅支持单节 DOCX；检测到多节文档，请先另存为单节副本。")
    section = document.sections[0]
    width_mm = section.page_width.mm
    height_mm = section.page_height.mm
    if width_mm >= height_mm or abs(width_mm - 210) > 3 or abs(height_mm - 297) > 3:
        raise FormatError("POC 当前仅支持 A4 纵向页面（约 210 × 297 mm）。")
    root = _xml(parts[MAIN])
    rel_root = _xml(parts.get("word/_rels/document.xml.rels", b"<Relationships/>"))
    external = sum(1 for rel in rel_root if rel.get("TargetMode") == "External")
    return {
        "sections": 1,
        "paragraphs": len(_paragraphs(root)),
        "tables": len(root.findall(".//w:tbl", NS)),
        "images": sum(1 for name in parts if name.startswith("word/media/")),
        "headers": sum(1 for name in parts if re.fullmatch(r"word/header\d+\.xml", name)),
        "footers": sum(1 for name in parts if re.fullmatch(r"word/footer\d+\.xml", name)),
        "external_hyperlinks": external,
        "page": f"A4 纵向 {width_mm:.0f} × {height_mm:.0f} mm",
        "size_bytes": len(data),
    }


def build_structure_mapping(data: bytes) -> list[dict]:
    """Return stable paragraph positions, recognized roles and weak role candidates."""
    rows = inspect_document(data)
    nonempty = [row["index"] for row in rows if row["text"].strip()]
    tail = set(nonempty[-3:])
    for row in rows:
        text = row["text"].strip()
        candidate = row["scope"]
        needs_confirmation = False
        if row["scope"] == "body" and text:
            if row["index"] in tail and (
                re.search(r"\d{4}\s*年\s*\d{1,2}\s*月", text)
                or re.fullmatch(r".{2,24}(?:处|室|局|校|院|部|委员会|中心)", text)
            ):
                candidate = "signature"
                needs_confirmation = True
            elif len(text) <= 30 and re.match(r"^(?:第[一二三四五六七八九十]+[章节]|[一二三四五六七八九十]+[、.]|（[一二三四五六七八九十]+）|\d+[.、])", text):
                candidate = "heading_1"
                needs_confirmation = True
        row.update(
            {
                "recognized_scope": row["scope"],
                "candidate_scope": candidate,
                "scope": candidate,
                "needs_confirmation": needs_confirmation,
                "excluded_reason": "",
            }
        )
    return rows


def _style_chain(style):
    seen = set()
    while style is not None and style.style_id not in seen:
        seen.add(style.style_id)
        yield style
        style = style.base_style


def _font_name_from_rpr(rpr):
    if rpr is None:
        return None
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        return None
    return fonts.get(qn("w:eastAsia")) or fonts.get(qn("w:ascii")) or fonts.get(qn("w:hAnsi"))


def _effective_run_value(run: Run, paragraph: Paragraph, prop: str):
    direct = run.font
    if prop == "font_family":
        value = _font_name_from_rpr(run._r.rPr)
    elif prop == "font_size":
        value = direct.size.pt if direct.size is not None else None
    elif prop == "bold":
        value = direct.bold
    elif prop == "italic":
        value = direct.italic
    elif prop == "color":
        value = f"#{direct.color.rgb}" if direct.color.rgb is not None else None
    else:
        raise FormatError(f"不支持的字符属性：{prop}")
    if value is not None:
        return value

    styles = []
    if run.style is not None:
        styles.extend(_style_chain(run.style))
    if paragraph.style is not None:
        styles.extend(_style_chain(paragraph.style))
    for style in styles:
        font = style.font
        if prop == "font_family":
            value = _font_name_from_rpr(style.element.rPr)
        elif prop == "font_size":
            value = font.size.pt if font.size is not None else None
        elif prop == "bold":
            value = font.bold
        elif prop == "italic":
            value = font.italic
        else:
            value = f"#{font.color.rgb}" if font.color.rgb is not None else None
        if value is not None:
            return value
    return None


def _length_pt(value):
    return value.pt if value is not None else None


def _line_spacing_value(value, rule=None):
    if value is None:
        return None
    if hasattr(value, "pt"):
        mode = "at_least" if rule == WD_LINE_SPACING.AT_LEAST else "exact"
        return {"mode": mode, "value": round(value.pt, 4)}
    return {"mode": "multiple", "value": round(float(value), 4)}


def _effective_paragraph_value(paragraph: Paragraph, prop: str):
    formats = [paragraph.paragraph_format]
    if paragraph.style is not None:
        formats.extend(style.paragraph_format for style in _style_chain(paragraph.style))
    for fmt in formats:
        if prop == "alignment":
            raw = fmt.alignment
            value = next((key for key, enum in ALIGNMENTS.items() if raw == enum), None)
        elif prop == "line_spacing":
            value = _line_spacing_value(fmt.line_spacing, fmt.line_spacing_rule)
        elif prop == "keep_with_next":
            value = fmt.keep_with_next
        elif prop in {"first_line_indent", "left_indent", "right_indent", "space_before", "space_after"}:
            value = _length_pt(getattr(fmt, prop))
        else:
            raise FormatError(f"不支持的段落属性：{prop}")
        if value is not None:
            return value
    return None


def _canonical(value):
    if isinstance(value, float):
        return round(value, 4)
    if isinstance(value, dict):
        return {key: _canonical(item) for key, item in value.items()}
    return value


def _same_value(current, target):
    return _canonical(current) == _canonical(target)


def _set_run_value(run: Run, prop: str, target):
    if prop == "font_family":
        fonts = run._r.get_or_add_rPr().get_or_add_rFonts()
        for attr in FONT_ATTRS:
            fonts.set(qn("w:" + attr), str(target))
        for attr in FONT_THEMES:
            fonts.attrib.pop(qn("w:" + attr), None)
    elif prop == "font_size":
        run.font.size = Pt(float(target))
        rpr = run._r.get_or_add_rPr()
        size_cs = rpr.find(qn("w:szCs"))
        if size_cs is None:
            size_cs = OxmlElement("w:szCs")
            anchor = rpr.find(qn("w:sz"))
            rpr.append(size_cs) if anchor is None else anchor.addnext(size_cs)
        size_cs.set(qn("w:val"), str(round(float(target) * 2)))
    elif prop == "color":
        run.font.color.rgb = RGBColor.from_string(str(target).lstrip("#").upper())
        color = run._r.find("w:rPr/w:color", NS)
        for attr in ("themeColor", "themeTint", "themeShade"):
            color.attrib.pop(qn("w:" + attr), None)
    elif prop == "bold":
        run.font.bold = bool(target)
    elif prop == "italic":
        run.font.italic = bool(target)


def _set_paragraph_value(paragraph: Paragraph, prop: str, target):
    fmt = paragraph.paragraph_format
    if prop == "alignment":
        fmt.alignment = ALIGNMENTS[target]
    elif prop == "line_spacing":
        if target["mode"] == "exact":
            fmt.line_spacing = Pt(target["value"])
            fmt.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        else:
            fmt.line_spacing = float(target["value"])
            fmt.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    elif prop == "keep_with_next":
        fmt.keep_with_next = bool(target)
    elif prop in {"first_line_indent", "left_indent", "right_indent", "space_before", "space_after"}:
        setattr(fmt, prop, Pt(float(target)))


def _rule_target(rule: dict):
    target = rule.get("normalized_target", rule.get("target"))
    if rule["property"] in {"font_size", "first_line_indent", "left_indent", "right_indent", "space_before", "space_after", *PAGE_PROPERTIES}:
        return float(target)
    if rule["property"] in {"bold", "italic", "keep_with_next"}:
        return bool(target)
    return target


def current_values(data: bytes, mapping: list[dict], scope: str, prop: str) -> list:
    """Return distinct effective values in stable display order."""
    if prop not in SUPPORTED_PROPERTIES:
        return []
    document = Document(BytesIO(data))
    if prop in PAGE_PROPERTIES:
        section = document.sections[0]
        return [round(getattr(section, prop).pt, 4)]
    roles = {int(row["index"]): row["scope"] for row in mapping}
    values = []
    for index, element in enumerate(_paragraphs(document.element), 1):
        role = roles.get(index, "body")
        if role == "preserve" or (scope != "all" and role != scope):
            continue
        paragraph = Paragraph(element, document)
        if prop in CHARACTER_PROPERTIES:
            for run_element in _runs(element):
                value = _effective_run_value(Run(run_element, document), paragraph, prop)
                if value not in values:
                    values.append(value)
        else:
            value = _effective_paragraph_value(paragraph, prop)
            if value not in values:
                values.append(value)
    return values


def _strip_allowed(root, rules: list[dict], mapping: list[dict]):
    root = deepcopy(root)
    roles = {int(row["index"]): row["scope"] for row in mapping}
    active = [rule for rule in rules if rule.get("apply", True)]
    for index, paragraph in enumerate(_paragraphs(root), 1):
        role = roles.get(index, "body")
        props = {rule["property"] for rule in active if rule["scope"] in {"all", role}}
        if role == "preserve" or not props:
            continue
        for run in _runs(paragraph):
            rpr = run.find("w:rPr", NS)
            if rpr is None:
                continue
            tag_map = {"font_size": ("sz", "szCs"), "color": ("color",), "bold": ("b", "bCs"), "italic": ("i", "iCs")}
            for prop, tags in tag_map.items():
                if prop in props:
                    for tag in tags:
                        for child in rpr.findall("w:" + tag, NS):
                            rpr.remove(child)
            fonts = rpr.find("w:rFonts", NS)
            if fonts is not None and "font_family" in props:
                for attr in FONT_ATTRS + FONT_THEMES:
                    fonts.attrib.pop(qn("w:" + attr), None)
                if not len(fonts) and not fonts.attrib:
                    rpr.remove(fonts)
            if not len(rpr) and not rpr.attrib:
                run.remove(rpr)
        ppr = paragraph.find("w:pPr", NS)
        if ppr is not None:
            simple = {"alignment": "jc", "keep_with_next": "keepNext"}
            for prop, tag in simple.items():
                if prop in props:
                    for child in ppr.findall("w:" + tag, NS):
                        ppr.remove(child)
            ind = ppr.find("w:ind", NS)
            if ind is not None:
                attrs = []
                if "first_line_indent" in props:
                    attrs += ["firstLine", "firstLineChars", "hanging", "hangingChars"]
                if "left_indent" in props:
                    attrs += ["left", "leftChars", "start", "startChars"]
                if "right_indent" in props:
                    attrs += ["right", "rightChars", "end", "endChars"]
                for attr in attrs:
                    ind.attrib.pop(qn("w:" + attr), None)
                if not len(ind) and not ind.attrib:
                    ppr.remove(ind)
            spacing = ppr.find("w:spacing", NS)
            if spacing is not None:
                attrs = []
                if "line_spacing" in props:
                    attrs += ["line", "lineRule"]
                if "space_before" in props:
                    attrs += ["before", "beforeLines", "beforeAutospacing"]
                if "space_after" in props:
                    attrs += ["after", "afterLines", "afterAutospacing"]
                for attr in attrs:
                    spacing.attrib.pop(qn("w:" + attr), None)
                if not len(spacing) and not spacing.attrib:
                    ppr.remove(spacing)
            if not len(ppr) and not ppr.attrib:
                paragraph.remove(ppr)
    page_props = {rule["property"] for rule in active if rule["scope"] == "page"}
    for section in root.findall(".//w:sectPr", NS):
        margin = section.find("w:pgMar", NS)
        if margin is not None:
            for prop in page_props:
                margin.attrib.pop(qn("w:" + prop.removesuffix("_margin")), None)
    return etree.tostring(root, method="c14n")


def validate_confirmed_result(source: bytes, output: bytes, rules: list[dict], mapping: list[dict]):
    before, after = _package(source), _package(output)
    if before.keys() != after.keys() or any(before[name] != after[name] for name in before if name != MAIN):
        raise FormatError("校验失败：文档附件、关系或其他部件发生了变化。")
    if _strip_allowed(_xml(before[MAIN]), rules, mapping) != _strip_allowed(_xml(after[MAIN]), rules, mapping):
        raise FormatError("校验失败：正文、对象或未选择的属性发生了变化。")
    document = Document(BytesIO(output))
    roles = {int(row["index"]): row["scope"] for row in mapping}
    active = [rule for rule in rules if rule.get("apply", True)]
    for index, element in enumerate(_paragraphs(document.element), 1):
        role = roles.get(index, "body")
        if role == "preserve":
            continue
        paragraph = Paragraph(element, document)
        for rule in active:
            if rule["scope"] not in {"all", role}:
                continue
            prop, target = rule["property"], _rule_target(rule)
            values = (
                [_effective_run_value(Run(run, document), paragraph, prop) for run in _runs(element)]
                if prop in CHARACTER_PROPERTIES
                else [_effective_paragraph_value(paragraph, prop)]
            )
            if values and any(not _same_value(value, target) for value in values):
                raise FormatError(f"校验失败：第 {index} 段的{PROPERTY_LABELS[prop]}未达到确认值。")
    for rule in active:
        if rule["scope"] == "page":
            actual = getattr(document.sections[0], rule["property"]).pt
            if not _same_value(actual, _rule_target(rule)):
                raise FormatError(f"校验失败：{PROPERTY_LABELS[rule['property']]}未达到确认值。")


def apply_confirmed_rules(data: bytes, rules: list[dict], mapping: list[dict]):
    """Apply one immutable, conflict-free rule snapshot to a working copy."""
    preflight_document(data)
    active = [rule for rule in rules if rule.get("apply", True)]
    if not active:
        raise FormatError("请至少选择一条完整且无冲突的要求。")
    unknown = {rule.get("property") for rule in active} - SUPPORTED_PROPERTIES
    if unknown:
        raise FormatError("存在当前引擎不支持的要求：" + "、".join(sorted(str(item) for item in unknown)))
    incompatible = [
        rule
        for rule in active
        if (rule.get("scope") == "page") != (rule.get("property") in PAGE_PROPERTIES)
    ]
    if incompatible:
        raise FormatError("页面规则只能使用页边距属性，页边距属性也只能作用于页面。")
    roles = {int(row["index"]): row["scope"] for row in mapping}
    document = Document(BytesIO(data))
    details = []
    changed_runs = set()
    changed_paragraphs = set()

    for rule in active:
        prop, target, scope = rule["property"], _rule_target(rule), rule["scope"]
        affected = []
        before_values = []
        if scope == "page":
            section = document.sections[0]
            current = getattr(section, prop).pt
            before_values.append(current)
            if not _same_value(current, target):
                setattr(section, prop, Pt(float(target)))
                affected.append("页面")
        else:
            for index, element in enumerate(_paragraphs(document.element), 1):
                role = roles.get(index, "body")
                if role == "preserve" or scope not in {"all", role}:
                    continue
                paragraph = Paragraph(element, document)
                if prop in CHARACTER_PROPERTIES:
                    changed_here = False
                    for run_element in _runs(element):
                        run = Run(run_element, document)
                        current = _effective_run_value(run, paragraph, prop)
                        if current not in before_values:
                            before_values.append(current)
                        if not _same_value(current, target):
                            _set_run_value(run, prop, target)
                            changed_runs.add((index, id(run_element)))
                            changed_here = True
                    if changed_here:
                        affected.append(index)
                else:
                    current = _effective_paragraph_value(paragraph, prop)
                    if current not in before_values:
                        before_values.append(current)
                    if not _same_value(current, target):
                        _set_paragraph_value(paragraph, prop, target)
                        changed_paragraphs.add(index)
                        affected.append(index)
        details.append(
            {
                "rule_id": rule.get("id"),
                "scope": scope,
                "property": prop,
                "before": before_values,
                "target": target,
                "affected": affected,
                "changed_count": len(affected),
            }
        )

    if not any(item["changed_count"] for item in details):
        return data, {"changed": False, "changed_runs": 0, "changed_paragraphs": 0, "details": details}
    output = BytesIO()
    with ZipFile(BytesIO(data)) as source, ZipFile(output, "w") as destination:
        destination.comment = source.comment
        for entry in source.infolist():
            raw = etree.tostring(document.element, xml_declaration=True, encoding="UTF-8", standalone=True) if entry.filename == MAIN else source.read(entry.filename)
            destination.writestr(entry, raw)
    result = output.getvalue()
    validate_confirmed_result(data, result, active, mapping)
    return result, {
        "changed": True,
        "changed_runs": len(changed_runs),
        "changed_paragraphs": len(changed_paragraphs),
        "details": details,
    }


def safe_stem(name: str) -> tuple[str, int]:
    name = PurePosixPath(name.replace("\\", "/")).name
    stem = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", name.rsplit(".", 1)[0]).strip(" .") or "文档"
    match = re.search(r"_v(\d+)$", stem)
    start = int(match.group(1)) + 1 if match else 1
    if match:
        stem = stem[: match.start()]
    return stem[:100], start


def snapshot_json(rules: list[dict], mapping: list[dict]) -> str:
    return json.dumps({"rules": rules, "mapping": mapping}, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
