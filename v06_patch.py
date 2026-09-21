"""Limited OOXML edits plus an independent character-level preservation verifier."""
from __future__ import annotations

from copy import deepcopy
from io import BytesIO
from zipfile import ZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.text.paragraph import Paragraph
from docx.text.run import Run
from lxml import etree

from docx_formatting import FormatError, MAIN, NS, _paragraphs, _xml
from prd_formatting import PAGE_PROPERTIES, PARAGRAPH_PROPERTIES, _effective_paragraph_value, _set_paragraph_value, _set_run_value, _style_chain
from v06_model import MODEL_VERSION, TEXT_TAGS, build_model, digest, node_text, run_spans
from v06_plan import CHAR_PROPS, EditPlan, encode

ENGINE_VERSION = "v06-patch-1"
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"
FONT_FIELDS = {"font_east_asia": ("eastAsia",), "font_western": ("ascii", "hAnsi", "cs")}
FONT_THEMES = {"font_east_asia": ("eastAsiaTheme",), "font_western": ("asciiTheme", "hAnsiTheme", "cstheme", "csTheme")}


def _rpr_chain(run, paragraph, document):
    yield run.find(qn("w:rPr"))
    wrapped = Run(run, document)
    if wrapped.style is not None:
        yield from (s.element.rPr for s in _style_chain(wrapped.style))
    if paragraph.style is not None:
        yield from (s.element.rPr for s in _style_chain(paragraph.style))
    yield document.styles.element.find("w:docDefaults/w:rPrDefault/w:rPr", NS)


def effective_run(run, paragraph, prop, document):
    chain = list(_rpr_chain(run, paragraph, document))
    if prop in FONT_FIELDS:
        values = []
        for attr in FONT_FIELDS[prop]:
            found = None
            for rpr in chain:
                f = rpr.find("w:rFonts", NS) if rpr is not None else None
                if f is not None:
                    theme = "cstheme" if attr == "cs" else attr + "Theme"
                    if f.get(qn("w:" + theme)) is not None:
                        found = "theme:" + f.get(qn("w:" + theme))
                        break
                    if f.get(qn("w:" + attr)) is not None:
                        found = f.get(qn("w:" + attr))
                        break
            values.append(found)
        return values[0] if len(set(values)) == 1 else values
    tags = {"font_size": "sz", "bold": "b", "italic": "i", "underline": "u", "color": "color"}
    tag = tags[prop]
    if prop == "font_size":
        sizes = []
        for name in ("sz", "szCs"):
            value = None
            for rpr in chain:
                n = rpr.find("w:" + name, NS) if rpr is not None else None
                if n is not None and n.get(qn("w:val")) is not None:
                    value = float(n.get(qn("w:val"))) / 2
                    break
            sizes.append(value)
        # Missing complex-script size follows ordinary size. A differing size
        # must not be treated as an already-met target.
        if sizes[1] is None:
            sizes[1] = sizes[0]
        return sizes[0] if sizes[0] == sizes[1] else sizes
    if prop in {"bold", "italic"}:
        direct = chain[0].find("w:" + tag, NS) if chain[0] is not None else None
        truth = lambda n: n.get(qn("w:val"), "1") not in {"0", "false", "off"}
        if direct is not None:
            return truth(direct)
        # Style booleans are toggle properties (ECMA-376); defaults are absolute.
        default = chain[-1].find("w:" + tag, NS) if chain[-1] is not None else None
        current = truth(default) if default is not None else False
        for rpr in reversed(chain[1:-1]):
            node = rpr.find("w:" + tag, NS) if rpr is not None else None
            if node is not None and truth(node):
                current = not current
        return current
    for rpr in chain:
        node = rpr.find("w:" + tag, NS) if rpr is not None else None
        if node is None:
            continue
        value = node.get(qn("w:val"))
        if prop == "font_size":
            return float(value) / 2 if value is not None else None
        if prop == "underline":
            return False if value in {"none", "0", "false"} else True if value == "single" else value
        if prop == "color":
            return "theme:" + node.get(qn("w:themeColor")) if node.get(qn("w:themeColor")) else "#" + value.upper() if value and value != "auto" else "auto"
    return False if prop == "underline" else None


def _clean_empty(parent, child):
    if child is not None and not len(child) and not child.attrib:
        parent.remove(child)


def _mask_rpr(run, props):
    rpr = run.find(qn("w:rPr"))
    if rpr is None:
        return
    for prop, tags in {"font_size": ("sz", "szCs"), "bold": ("b",), "italic": ("i",), "underline": ("u",), "color": ("color",)}.items():
        if prop in props:
            for tag in tags:
                for node in rpr.findall("w:" + tag, NS):
                    rpr.remove(node)
    fonts = rpr.find("w:rFonts", NS)
    if fonts is not None:
        for prop in props & FONT_FIELDS.keys():
            for attr in FONT_FIELDS[prop] + FONT_THEMES[prop]:
                fonts.attrib.pop(qn("w:" + attr), None)
        _clean_empty(rpr, fonts)
    _clean_empty(run, rpr)


def _mask_ppr(p, props):
    ppr = p.find("w:pPr", NS)
    if ppr is None:
        return
    for prop, tag in (("alignment", "jc"), ("keep_with_next", "keepNext")):
        if prop in props:
            for node in ppr.findall("w:" + tag, NS):
                ppr.remove(node)
    for tag, mapping in (("ind", {"first_line_indent": ("firstLine", "firstLineChars", "hanging", "hangingChars"), "left_indent": ("left", "leftChars", "start", "startChars"), "right_indent": ("right", "rightChars", "end", "endChars")}), ("spacing", {"line_spacing": ("line", "lineRule"), "space_before": ("before", "beforeLines", "beforeAutospacing"), "space_after": ("after", "afterLines", "afterAutospacing")})):
        node = ppr.find("w:" + tag, NS)
        if node is not None:
            for prop in props & mapping.keys():
                for attr in mapping[prop]:
                    node.attrib.pop(qn("w:" + attr), None)
            _clean_empty(ppr, node)
    _clean_empty(p, ppr)


def _set_character(run, prop, value, document):
    wrapped = Run(run, document)
    if prop in FONT_FIELDS:
        fonts = wrapped._r.get_or_add_rPr().get_or_add_rFonts()
        for attr in FONT_FIELDS[prop]:
            fonts.set(qn("w:" + attr), value)
        for attr in FONT_THEMES[prop]:
            fonts.attrib.pop(qn("w:" + attr), None)
    elif prop == "underline":
        wrapped.font.underline = value
    else:
        _set_run_value(wrapped, prop, value)


def _attributes_at(assignments, pid, offset):
    return {a["property"]: a["value"] for a in assignments if a["pid"] == pid and a["property"] in CHAR_PROPS and a["start"] <= offset < a["end"]}


def _check_plan(model, plan):
    data = plan.read()
    if data["input_hash"] != model.source_hash or data["input_revision"] != model.revision or data["model_version"] != MODEL_VERSION:
        raise FormatError("输入版本或结构映射已变化，必须重新生成计划。")
    if data["roles"] != {str(k): v for k, v in model.role_overrides.items()}:
        raise FormatError("角色映射已变化，请重新生成计划。")
    # Recompute the unique assignments from the frozen source intentions. A caller
    # cannot manufacture extra allowed writes by editing a plan payload.
    from v06_plan import build_plan
    rebuilt = build_plan(model, data["intents"], data["protections"])
    if not rebuilt.plan or rebuilt.plan.payload != plan.payload:
        raise FormatError("计划与其冻结来源不一致。")
    return data


def current_values(model, scope, prop):
    from v06_plan import resolve_scope
    spans, _ = resolve_scope(model, scope)
    document = Document(BytesIO(model.data))
    if prop in PAGE_PROPERTIES:
        return [getattr(document.sections[0], prop).pt]
    elements = _paragraphs(document.element)
    values = []
    for span in spans:
        p = elements[span["pid"]-1]
        paragraph = Paragraph(p, document)
        if prop in PARAGRAPH_PROPERTIES:
            observed = [_effective_paragraph_value(paragraph, prop)]
        else:
            observed = [effective_run(r, paragraph, prop, document) for r, a, b, _, protected in run_spans(p) if not protected and a < b and a < span["end"] and b > span["start"]]
        for value in observed:
            if value not in values:
                values.append(value)
    return values


def patch_document(model, plan):
    frozen = _check_plan(model, plan)
    assignments = frozen["assignments"]
    document = Document(BytesIO(model.data))
    changed = []
    for pid, p in enumerate(_paragraphs(document.element), 1):
        paragraph = Paragraph(p, document)
        relevant = [a for a in assignments if a["pid"] == pid]
        if not relevant:
            continue
        for assignment in relevant:
            prop, value = assignment["property"], assignment["value"]
            if prop in PARAGRAPH_PROPERTIES and _effective_paragraph_value(paragraph, prop) != value:
                # Remove only mutually exclusive representations of THIS property.
                _mask_ppr(p, {prop})
                _set_paragraph_value(paragraph, prop, value)
                changed.append(dict(assignment))
        for run, a, b, _, protected in run_spans(p):
            if protected or a == b:
                continue
            # Each piece retains the entire original rPr and original run attrs.
            # Non-text children receive no character edit and keep their position.
            groups = []
            offset = a
            for child in run:
                if child.tag == qn("w:rPr"):
                    continue
                if child.tag in TEXT_TAGS:
                    text = node_text(child)
                    boundaries = sorted({offset, offset + len(text)} | {x[k] for x in relevant if x["property"] in CHAR_PROPS for k in ("start", "end") if offset < x[k] < offset + len(text)})
                    segments = list(zip(boundaries, boundaries[1:])) or [(offset, offset)]
                    for x, y in segments:
                        requested = _attributes_at(relevant, pid, x) if x < y else {}
                        writes = {k: v for k, v in requested.items() if effective_run(run, paragraph, k, document) != v}
                        segment = deepcopy(child)
                        if child.tag == qn("w:t"):
                            segment.text = text[x-offset:y-offset]
                            if segment.text and (segment.text[0].isspace() or segment.text[-1].isspace()):
                                segment.set(XML_SPACE, "preserve")
                        if groups and groups[-1][0] == writes:
                            groups[-1][1].append(segment)
                        else:
                            groups.append((writes, [segment]))
                        if writes:
                            changed.append({"pid": pid, "start": x, "end": y, "set": writes})
                    offset += len(text)
                else:
                    if groups and not groups[-1][0]:
                        groups[-1][1].append(deepcopy(child))
                    else:
                        groups.append(({}, [deepcopy(child)]))
            if not any(properties for properties, _ in groups):
                continue
            parent = run.getparent()
            index = parent.index(run)
            for properties, children in groups:
                new = deepcopy(run)
                for child in list(new):
                    if child.tag != qn("w:rPr"):
                        new.remove(child)
                new.extend(children)
                for prop, value in properties.items():
                    _set_character(new, prop, value, document)
                parent.insert(index, new)
                index += 1
            parent.remove(run)
    for a in assignments:
        if a["property"] in PAGE_PROPERTIES:
            section = document.sections[0]
            if abs(getattr(section, a["property"]).pt - a["value"]) > 0.05:
                setattr(section, a["property"], Pt(a["value"]))
                changed.append(dict(a))
    if not changed:
        return model.data, []
    output = BytesIO()
    with ZipFile(BytesIO(model.data)) as old, ZipFile(output, "w") as new:
        new.comment = old.comment
        for entry in old.infolist():
            new.writestr(entry, etree.tostring(document.element, xml_declaration=True, encoding="UTF-8", standalone=True) if entry.filename == MAIN else old.read(entry.filename))
    return output.getvalue(), changed


def _semantic_tree(raw, assignments):
    root = _xml(raw)
    for pid, p in enumerate(_paragraphs(root), 1):
        relevant = [a for a in assignments if a["pid"] == pid]
        _mask_ppr(p, {a["property"] for a in relevant if a["property"] in PARAGRAPH_PROPERTIES})
        for run, start, _, _, _ in run_spans(p):
            parent, index = run.getparent(), run.getparent().index(run)
            offset = start
            content = [c for c in run if c.tag != qn("w:rPr")]
            if not content:
                continue
            for child in content:
                if child.tag in TEXT_TAGS and node_text(child):
                    pieces = [(ch, offset + n) for n, ch in enumerate(node_text(child))]
                else:
                    pieces = [(None, None)]
                for ch, pos in pieces:
                    atom = deepcopy(run)
                    for c in list(atom):
                        if c.tag != qn("w:rPr"):
                            atom.remove(c)
                    node = deepcopy(child)
                    if node.tag == qn("w:t") and ch is not None:
                        node.text = ch
                        node.attrib.pop(XML_SPACE, None)
                    atom.append(node)
                    _mask_rpr(atom, set(_attributes_at(relevant, pid, pos)) if pos is not None else set())
                    # Empty rPr is equivalent to absent rPr after run splitting.
                    _clean_empty(atom, atom.find(qn("w:rPr")))
                    parent.insert(index, atom)
                    index += 1
                if child.tag in TEXT_TAGS:
                    offset += len(node_text(child))
            parent.remove(run)
    for a in assignments:
        if a["property"] in PAGE_PROPERTIES:
            for margin in root.findall(".//w:sectPr/w:pgMar", NS):
                margin.attrib.pop(qn("w:" + a["property"].removesuffix("_margin")), None)
    return etree.tostring(root, method="c14n")


def verify_result(model, output, plan):
    """Does not call the patcher: independently compares scope, properties and parts."""
    frozen = _check_plan(model, plan)
    result = build_model(output)
    if model.parts.keys() != result.parts.keys() or any(model.parts[n] != result.parts[n] for n in model.parts if n != MAIN):
        raise FormatError("保全失败：图片、批注、嵌入对象、关系或其他包部件发生变化。")
    assignments = frozen["assignments"]
    if _semantic_tree(model.parts[MAIN], assignments) != _semantic_tree(result.parts[MAIN], assignments):
        raise FormatError("保全失败：内容、对象、范围外或未指定属性发生变化。")
    document = Document(BytesIO(output))
    paragraphs = _paragraphs(document.element)
    for a in assignments:
        prop, target = a["property"], a["value"]
        if prop in PAGE_PROPERTIES:
            values = [getattr(document.sections[0], prop).pt]
        else:
            p = paragraphs[a["pid"]-1]
            paragraph = Paragraph(p, document)
            if prop in PARAGRAPH_PROPERTIES:
                values = [_effective_paragraph_value(paragraph, prop)]
            else:
                values = [effective_run(run, paragraph, prop, document) for run, x, y, _, _ in run_spans(p) if x < y and x < a["end"] and y > a["start"]]
        def same(value):
            if type(target) in {int, float} and type(value) in {int, float}:
                return abs(target-value) <= 0.051
            if isinstance(target, dict) and isinstance(value, dict):
                return value["mode"] == target["mode"] and abs(value["value"]-target["value"]) <= 0.005
            return value == target
        if not values or not all(same(v) for v in values):
            raise FormatError(f"目标检查失败：段落 {a['pid']} 的 {prop} 未达到计划值。")
    return {"structure": "passed", "properties": "passed", "preservation": "passed", "render": "not_run", "visual_review": "pending", "client_compatibility": "not_run", "inventory": result.inventory}
